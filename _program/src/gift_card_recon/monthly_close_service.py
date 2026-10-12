from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
import uuid
from zipfile import BadZipFile
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal
from functools import wraps
from pathlib import Path
from typing import Callable, Iterable, Mapping, Sequence

from gift_card_recon.close_assessment import (
    CloseAssessment,
    CloseStatus,
    ControlDisposition,
    ControlOutcome,
    assess_monthly_close,
    build_close_assessment,
    integrity_control,
)
from gift_card_recon.business_identity import (
    CorrectionRequired, correction_token, input_set_identity, workbook_values_identity,
)
from gift_card_recon.publication import (
    PublicationTransaction, file_state, recover_pending, support_root_for, writer_lock,
    retry_file_operation as _retry_transient_file_operation,
)
from gift_card_recon.darden import build_monthly_close_certification, parse_darden_credit_memo
from gift_card_recon.evidence_archive import (
    ArchiveError,
    ArchiveRecord,
    EvidenceItem,
    ManifestArtifact,
    plan_evidence_archive,
    write_close_manifest_atomic,
)
from gift_card_recon.excel_writer import write_reconciliation_workbook
from gift_card_recon.fiscal_calendar import FiscalPeriod, fiscal_period_for_label
from gift_card_recon.micros import (
    MicrosEvidence,
    build_weekly_pos_variances,
    load_micros_evidence,
    period_tender_variance,
    resolve_micros_export_dir,
    validate_micros_source,
    weekly_tender_variances,
)
from gift_card_recon.models import (
    DardenCreditMemo,
    PosControls,
    ReconciliationResult,
    SourceFileAudit,
    WeeklyPosVariance,
)
from gift_card_recon.monthly_report import (
    DEFAULT_EVIDENCE_LABELS,
    MonthlyCloseReportData,
    WeeklyCloseReportRow,
    write_monthly_close_report_workbook,
)
from gift_card_recon.monthly_explanations import (
    locate_monthly_variance_explanation_source,
    read_monthly_variance_explanations,
)
from gift_card_recon.parsers import ParseError, discover_input_files, parse_activity_file, parse_summary, summary_report_end
from gift_card_recon.pdf_export import PdfExportError, export_monthly_close_report_pdf
from gift_card_recon.reconcile import build_reconciliation
from gift_card_recon.source_validation import validate_activity_evidence
from gift_card_recon.store_config import REVIEW_VARIANCE_LIMIT, StoreConfig, get_store_config
from gift_card_recon.utils import file_modified_at, money, sha256_file
from gift_card_recon.variance_explanation import (
    EXPLANATION_INPUT_CELL,
    EXPLANATION_SHEET,
    WeeklyVarianceExplanation,
    read_variance_explanation_workbook,
    resolve_live_weekly_explanation_path,
    variance_control_mismatch_details,
    variance_explanation_path,
)


PdfExporter = Callable[..., Path]



class CloseBlockedError(RuntimeError):
    """A close was not published because one or more required controls failed."""

    def __init__(
        self,
        message: str,
        *,
        assessment: CloseAssessment,
        review_workbook: Path | None = None,
        review_pdf: Path | None = None,
        review_pdf_error: str | None = None,
        review_publication_error: str | None = None,
    ) -> None:
        super().__init__(message)
        self.assessment = assessment
        self.review_workbook = review_workbook
        self.review_pdf = review_pdf
        self.review_pdf_error = review_pdf_error
        self.review_publication_error = review_publication_error


@dataclass(frozen=True)
class ReviewDiagnosticResult:
    workbook_path: Path
    pdf_path: Path | None
    pdf_error: str | None = None

    def __iter__(self):
        """Keep the former workbook/PDF tuple-unpacking contract."""

        yield self.workbook_path
        yield self.pdf_path


@dataclass(frozen=True, slots=True)
class ArchivedVarianceExplanationSource:
    """One manifest-verified explanation source selected for archive reissue."""

    path: Path
    sha256: str
    size_bytes: int

    def __post_init__(self) -> None:
        if not isinstance(self.sha256, str):
            raise ValueError("Archived variance explanation SHA-256 is invalid.")
        digest = self.sha256.lower()
        if len(digest) != 64 or any(char not in "0123456789abcdef" for char in digest):
            raise ValueError("Archived variance explanation SHA-256 is invalid.")
        if (
            not isinstance(self.size_bytes, int)
            or isinstance(self.size_bytes, bool)
            or self.size_bytes < 0
        ):
            raise ValueError("Archived variance explanation size is invalid.")
        object.__setattr__(self, "path", Path(self.path))
        object.__setattr__(self, "sha256", digest)


@dataclass(frozen=True)
class MonthlyCloseRunResult:
    workbook_path: Path
    pdf_path: Path | None
    reconciliation: ReconciliationResult
    weekly_variances: tuple[WeeklyPosVariance, ...]
    assessment: CloseAssessment
    manifest_path: Path
    archive_records: tuple[ArchiveRecord, ...]
    input_set_id: str = ""
    result_id: str = ""
    disposition: str = "published"
    pdf_status: str = "complete"
    pdf_error: str | None = None
    recovery_actions: tuple[str, ...] = ()
    run_id: str = field(default_factory=lambda: uuid.uuid4().hex)

    def to_dict(self) -> dict:
        return {
            "restaurant": self.assessment.store, "store": self.assessment.store,
            "period": self.reconciliation.period, "input_set_id": self.input_set_id,
            "run_id": self.run_id, "result_id": self.result_id,
            "status": self.disposition, "business_disposition": self.assessment.status.value,
            "artifacts": {"workbook": str(self.workbook_path),
                          "pdf": str(self.pdf_path) if self.pdf_path else None,
                          "manifest": str(self.manifest_path)},
            "pdf_status": self.pdf_status, "pdf_error": self.pdf_error,
            "recovery_actions": list(self.recovery_actions),
        }

    def __iter__(self):
        """Compatibility with the former three-value result."""

        yield self.workbook_path
        yield self.reconciliation
        yield list(self.weekly_variances)


def canonical_output_paths(
    output_root: Path,
    *,
    config: StoreConfig,
    fiscal_period: FiscalPeriod,
) -> tuple[Path, Path]:
    folder = Path(output_root) / "Monthly Close" / fiscal_period.folder_name
    base = f"{config.output_slug}_{fiscal_period.period_key}_Monthly_Close"
    return folder / f"{base}.xlsx", folder / f"{base}.pdf"


def review_output_paths(
    output_root: Path,
    *,
    config: StoreConfig,
    fiscal_period: FiscalPeriod,
) -> tuple[Path, Path]:
    folder = Path(output_root) / "Monthly Close - Review Required"
    base = f"{config.output_slug}_{fiscal_period.period_key}_Review_Required"
    return folder / f"{base}.xlsx", folder / f"{base}.pdf"


def _publication_staging_directory(
    destination_dir: Path,
    *,
    prefix: str,
    support_root: Path | None = None,
) -> tempfile.TemporaryDirectory[str]:
    """Create a private workspace; the shared publisher stages on the target volume."""

    target_dir = Path(destination_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    staging_parent = Path(support_root) if support_root is not None else target_dir.parent
    staging_parent.mkdir(parents=True, exist_ok=True)
    workspace = tempfile.TemporaryDirectory(
        prefix=f".gift-card-{prefix}",
        dir=str(staging_parent),
    )
    try:
        staging_dir = Path(workspace.name)
        staging_dir.chmod(0o700)
    except Exception:
        workspace.cleanup()
        raise
    return workspace


def _resolve_monthly_store_config(
    store: str | int,
    supplied: StoreConfig | None,
) -> StoreConfig:
    canonical = get_store_config(store)
    if supplied is None:
        return canonical
    if replace(supplied, micros_default_path=canonical.micros_default_path) != canonical:
        raise ValueError(
            f"Supplied store configuration {supplied.store!r} does not match "
            f"requested store {canonical.store!r}."
        )
    return supplied


def _monthly_writer(function):
    @wraps(function)
    def locked(*args, **kwargs):
        support = support_root_for(Path(kwargs["archive_root"]))
        with writer_lock(support):
            recovered = tuple(recover_pending(support))
            result = function(*args, **kwargs)
            return replace(result, recovery_actions=(*recovered, *result.recovery_actions))
    return locked


@_monthly_writer
def run_monthly_close_service(
    *,
    store: str | int,
    store_config: StoreConfig | None = None,
    period: str,
    input_dir: Path,
    micros_path: Path | None,
    micros_work_dir: Path,
    archive_root: Path,
    output_root: Path,
    output_path: Path | None = None,
    darden_path: Path | None = None,
    darden_report: DardenCreditMemo | None = None,
    fiscal_period: FiscalPeriod | None = None,
    cleanup_sources: bool = True,
    allow_unconfigured_micros: bool = False,
    archived_variance_explanations: Mapping[
        date, ArchivedVarianceExplanationSource
    ] | None = None,
    archived_monthly_variance_explanation: ArchivedVarianceExplanationSource | None = None,
    archived_source_paths: Mapping[str, Sequence[Path]] | None = None,
    live_source_paths: Mapping[str, Sequence[Path]] | None = None,
    generated_at: datetime | None = None,
    pdf_exporter: PdfExporter | None = None,
    approve_correction: str | None = None,
) -> MonthlyCloseRunResult:
    """Assess, render, archive, publish, and clean up one store-period close."""

    config = _resolve_monthly_store_config(store, store_config)
    canonical_period = fiscal_period_for_label(period)
    if fiscal_period is not None and (
        fiscal_period.period_key != canonical_period.period_key
        or fiscal_period.folder_name != canonical_period.folder_name
        or fiscal_period.start_date != canonical_period.start_date
        or fiscal_period.end_date != canonical_period.end_date
    ):
        raise ParseError(f"Fiscal period {period!r} does not match the supplied period object.")
    fiscal_period = canonical_period
    generated_at = generated_at.astimezone() if generated_at is not None else datetime.now().astimezone()
    pdf_exporter = pdf_exporter or export_monthly_close_report_pdf
    input_dir = Path(input_dir)
    archive_root = Path(archive_root)
    output_root = Path(output_root)
    canonical_xlsx, canonical_pdf = canonical_output_paths(
        output_root,
        config=config,
        fiscal_period=fiscal_period,
    )
    if output_path is not None:
        canonical_xlsx = Path(output_path)
        if canonical_xlsx.suffix.lower() != ".xlsx":
            raise ParseError(f"Monthly-close workbook output must end in .xlsx: {canonical_xlsx}")
        canonical_pdf = canonical_xlsx.with_suffix(".pdf")
    manifest_path = archive_root / "Monthly Close" / config.store / fiscal_period.folder_name / "close_manifest.json"
    manifest_state = file_state(manifest_path)
    report_candidates = [canonical_xlsx, canonical_pdf]
    if canonical_xlsx.parent.is_dir():
        report_candidates.extend(path for path in canonical_xlsx.parent.iterdir()
                                 if path.is_file() and path.suffix.lower() in {".xlsx", ".pdf"})
    initial_output_states = {path.resolve(): file_state(path) for path in report_candidates}

    try:
        close_data = _build_close_data(
            config=config,
            fiscal_period=fiscal_period,
            input_dir=input_dir,
            micros_path=Path(micros_path) if micros_path is not None else config.micros_default_path,
            micros_work_dir=Path(micros_work_dir),
            archive_root=archive_root,
            darden_path=Path(darden_path) if darden_path is not None else None,
            darden_report=darden_report,
            cleanup_sources=cleanup_sources,
            allow_unconfigured_micros=allow_unconfigured_micros,
            archived_variance_explanations=archived_variance_explanations,
            archived_monthly_variance_explanation=archived_monthly_variance_explanation,
            archived_source_paths=archived_source_paths,
            live_source_paths=live_source_paths,
            output_root=output_root,
            canonical_xlsx=canonical_xlsx,
        )
    except (ParseError, ArchiveError, OSError, ValueError, RuntimeError) as exc:
        assessment = _failure_assessment(config.store, "evidence_validation", "Evidence validation", str(exc))
        try:
            review = _write_review_diagnostic(
                output_root=output_root,
                config=config,
                fiscal_period=fiscal_period,
                assessment=assessment,
                generated_at=generated_at,
                message=str(exc),
                pdf_exporter=pdf_exporter,
            )
        except (OSError, ValueError, RuntimeError) as diagnostic_exc:
            raise CloseBlockedError(
                str(exc),
                assessment=assessment,
                review_publication_error=str(diagnostic_exc),
            ) from exc
        raise CloseBlockedError(
            str(exc),
            assessment=assessment,
            review_workbook=review.workbook_path,
            review_pdf=review.pdf_path,
            review_pdf_error=review.pdf_error,
        ) from exc

    assessment = close_data.assessment
    if not assessment.can_publish_close:
        blockers = "; ".join(control.message for control in assessment.blockers)
        blocker_message = blockers or "One or more monthly-close controls require review."
        try:
            review = _write_detailed_review(
                output_root=output_root,
                config=config,
                fiscal_period=fiscal_period,
                close_data=close_data,
                generated_at=generated_at,
                pdf_exporter=pdf_exporter,
            )
        except (OSError, ValueError, RuntimeError) as diagnostic_exc:
            raise CloseBlockedError(
                blocker_message,
                assessment=assessment,
                review_publication_error=str(diagnostic_exc),
            ) from None
        raise CloseBlockedError(
            blocker_message,
            assessment=assessment,
            review_workbook=review.workbook_path,
            review_pdf=review.pdf_path,
            review_pdf_error=review.pdf_error,
        )

    canonical_xlsx.parent.mkdir(parents=True, exist_ok=True)
    support = support_root_for(archive_root)
    try:
        for publish_target in (canonical_xlsx, canonical_pdf, manifest_path):
            _assert_publishable(publish_target)
    except OSError as exc:
        raise CloseBlockedError(str(exc), assessment=_with_blocking_control(
            assessment, code="publication_integrity", label="Publication target",
            message=str(exc),
        )) from exc
    _require_file_state(manifest_path, manifest_state)
    previous = _read_previous_close(manifest_path, store=config.store, period=fiscal_period.period_key)
    if previous:
        canonical_xlsx = resolve_monthly_workbook(previous, canonical_xlsx)
        canonical_pdf = _resolve_monthly_pdf(previous, canonical_xlsx.with_suffix(".pdf"))
    output_states = {canonical_xlsx: initial_output_states.get(canonical_xlsx.resolve()),
                     canonical_pdf: initial_output_states.get(canonical_pdf.resolve()), manifest_path: manifest_state}
    previous_id = str(previous.get("result_id") or (sha256_file(manifest_path) if previous else ""))
    previous_input_id = previous.get("input_set_id")
    if previous and previous_input_id is None:
        previous_input_id = _legacy_monthly_identity(
            config=config, fiscal_period=fiscal_period, archive_root=archive_root,
            micros_work_dir=Path(micros_work_dir),
        )
    unchanged_report = (not previous.get("workbook_values_identity")
                        or workbook_values_identity(canonical_xlsx) == previous["workbook_values_identity"])
    if previous and previous_input_id == close_data.input_set_id and unchanged_report:
        _verify_existing_workbook(canonical_xlsx, previous)
        preserved_records = _previous_archive_records(previous, archive_root=archive_root)
        pdf_record = next((row for row in previous.get("artifacts", []) if row.get("role") == "pdf"), None)
        pdf_complete = (previous.get("pdf_status", "complete") == "complete" and pdf_record is not None
                        and canonical_pdf.is_file() and canonical_pdf.stat().st_size == pdf_record.get("size_bytes")
                        and sha256_file(canonical_pdf) == pdf_record.get("sha256"))
        return MonthlyCloseRunResult(
            workbook_path=canonical_xlsx,
            pdf_path=canonical_pdf if pdf_complete else None,
            reconciliation=close_data.reconciliation, weekly_variances=close_data.weekly_variances,
            assessment=assessment, manifest_path=manifest_path, archive_records=preserved_records,
            input_set_id=close_data.input_set_id, result_id=previous_id, disposition="no_change",
            pdf_status="complete" if pdf_complete else "pending",
            pdf_error=None if pdf_complete else previous.get("pdf_error") or "PDF is missing or no longer matches this result; retry export.",
        )
    draft_input = not previous and _is_review_explanation_input(canonical_xlsx, close_data)
    if previous or (canonical_xlsx.exists() and not draft_input):
        previous_id = previous_id or sha256_file(canonical_xlsx)
        expected = correction_token(previous_id, close_data.input_set_id)
        if approve_correction != expected:
            changes = _monthly_changed_facts(previous, close_data)
            raise CorrectionRequired(previous_id, close_data.input_set_id, changes)
        if previous:
            _verify_existing_workbook(canonical_xlsx, previous)
    result_id = hashlib.sha256(
        f"monthly:{config.store}:{fiscal_period.period_key}:{previous_id}:{close_data.input_set_id}".encode()
    ).hexdigest()
    pdf_error = None
    cleanup_messages = []
    try:
        support.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="monthly-close-", dir=support) as temp_dir:
            temp_root = Path(temp_dir)
            temp_xlsx = temp_root / canonical_xlsx.name
            temp_pdf = temp_root / canonical_pdf.name
            archived_result = _with_archive_source_audits(close_data.reconciliation, close_data.archive_records)
            _write_full_workbook(
                output_path=temp_xlsx, result=archived_result, close_data=close_data,
                generated_at=generated_at, archive_published=True,
            )
            try:
                pdf_exporter(workbook_path=temp_xlsx, pdf_path=temp_pdf, expected_location_label=config.report_heading)
                if not temp_pdf.is_file() or temp_pdf.stat().st_size == 0:
                    raise PdfExportError("PDF export produced no usable file.")
            except (PdfExportError, OSError, ValueError, RuntimeError) as exc:
                pdf_error = str(exc)
                temp_pdf.unlink(missing_ok=True)
            with PublicationTransaction(support) as transaction:
                # Archive the exact captured evidence before replacing any result.
                for index, record in enumerate(close_data.archive_records):
                    if sha256_file(record.source_path) != record.sha256 or record.source_path.stat().st_size != record.size_bytes:
                        raise ArchiveError(f"Evidence changed before publication: {record.source_path}")
                    if record.archive_path.exists():
                        if sha256_file(record.archive_path) != record.sha256:
                            raise ArchiveError(f"Archived evidence changed: {record.archive_path}")
                        continue
                    staged_source = temp_root / f"evidence-{index}"
                    shutil.copy2(record.source_path, staged_source)
                    if sha256_file(staged_source) != record.sha256:
                        raise ArchiveError(f"Evidence copy did not verify: {record.source_path}")
                    transaction.replace(staged_source, record.archive_path, expected_before=None)
                revision_dir = manifest_path.parent / "revisions" / result_id[:16]
                if previous_id:
                    previous_dir = manifest_path.parent / "revisions" / previous_id[:16]
                    for index, old in enumerate((canonical_xlsx, canonical_pdf, manifest_path)):
                        retained_name = ("workbook.xlsx", "report.pdf", "close_manifest.json")[index]
                        if old.is_file() and not (previous_dir / retained_name).exists():
                            old_copy = temp_root / f"previous-{index}"
                            shutil.copy2(old, old_copy)
                            _require_file_state(old, output_states[old])
                            _require_file_state(old_copy, output_states[old])
                            transaction.replace(old_copy, previous_dir / retained_name, expected_before=None)
                workbook_values = workbook_values_identity(temp_xlsx)
                workbook_controls = _monthly_report_control_identity(temp_xlsx)
                baseline = temp_root / "revision-workbook.xlsx"
                shutil.copy2(temp_xlsx, baseline)
                transaction.replace(baseline, revision_dir / "workbook.xlsx", expected_before=None)
                workbook_artifact = _attest_staged_artifact(temp_xlsx, published_path=canonical_xlsx)
                transaction.replace(temp_xlsx, canonical_xlsx, expected_before=output_states[canonical_xlsx])
                artifacts = {"workbook": workbook_artifact}
                if pdf_error is None:
                    pdf_copy = temp_root / "revision-report.pdf"
                    shutil.copy2(temp_pdf, pdf_copy)
                    transaction.replace(pdf_copy, revision_dir / "report.pdf", expected_before=None)
                    artifacts["pdf"] = _attest_staged_artifact(temp_pdf, published_path=canonical_pdf)
                    transaction.replace(temp_pdf, canonical_pdf, expected_before=output_states[canonical_pdf])
                elif output_states[canonical_pdf] is not None or canonical_pdf.exists():
                    transaction.remove(canonical_pdf, expected_before=output_states[canonical_pdf])
                staged_manifest = temp_root / "close_manifest.json"
                write_close_manifest_atomic(
                    staged_manifest, store=config.store, location=config.location_name,
                    period=fiscal_period.period_key, status=assessment.status.value,
                    source_records=close_data.archive_records, artifacts=artifacts,
                    archive_root=archive_root, generated_at=generated_at,
                )
                payload = json.loads(staged_manifest.read_text(encoding="utf-8"))
                payload.update({
                    "schema_version": 2, "input_set_id": close_data.input_set_id,
                    "result_id": result_id, "previous_result_id": previous_id or None,
                    "pdf_status": "pending" if pdf_error else "complete", "pdf_error": pdf_error,
                    "workbook_values_identity": workbook_values,
                    "workbook_control_identity": workbook_controls,
                    "workbook_relative_path": (str(canonical_xlsx.resolve().relative_to(archive_root.parent.resolve()))
                                               if _is_within(canonical_xlsx, archive_root.parent) else None),
                    "business_facts": _monthly_business_facts(close_data),
                })
                staged_manifest.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
                manifest_copy = temp_root / "revision-manifest.json"
                shutil.copy2(staged_manifest, manifest_copy)
                transaction.replace(manifest_copy, revision_dir / "close_manifest.json", expected_before=None)
                transaction.replace(staged_manifest, manifest_path, expected_before=manifest_state)
                for artifact in artifacts.values():
                    _verify_published_artifact(artifact)
                cleanup = [record.source_path for record in close_data.archive_records
                           if cleanup_sources and record.remove_after_publish
                           and record.source_path.resolve() != record.archive_path.resolve()]
                cleanup_messages = transaction.commit(cleanup=cleanup, expected_cleanup={record.source_path: record.sha256 for record in close_data.archive_records if record.source_path in cleanup})
    except (ArchiveError, PdfExportError, OSError, ValueError, RuntimeError) as exc:
        failed = _with_blocking_control(
            assessment,
            code="publication_integrity",
            label="Publication and archive transaction",
            message=str(exc),
        )
        failed_data = replace(close_data, assessment=failed)
        try:
            review = _write_detailed_review(
                output_root=output_root,
                config=config,
                fiscal_period=fiscal_period,
                close_data=failed_data,
                generated_at=generated_at,
                pdf_exporter=pdf_exporter,
            )
        except (OSError, ValueError, RuntimeError) as diagnostic_exc:
            raise CloseBlockedError(
                str(exc),
                assessment=failed,
                review_publication_error=str(diagnostic_exc),
            ) from exc
        raise CloseBlockedError(
            str(exc),
            assessment=failed,
            review_workbook=review.workbook_path,
            review_pdf=review.pdf_path,
            review_pdf_error=review.pdf_error,
        ) from exc

    cleanup_messages.extend(_archive_stale_review_artifacts(
        output_root=output_root, archive_root=archive_root,
        config=config, fiscal_period=fiscal_period,
    ))
    return MonthlyCloseRunResult(
        workbook_path=canonical_xlsx,
        pdf_path=canonical_pdf if pdf_error is None else None,
        reconciliation=archived_result,
        weekly_variances=close_data.weekly_variances,
        assessment=assessment,
        manifest_path=manifest_path,
        archive_records=close_data.archive_records,
        input_set_id=close_data.input_set_id, result_id=result_id,
        pdf_status="pending" if pdf_error else "complete", pdf_error=pdf_error,
        recovery_actions=tuple(cleanup_messages),
    )


def _require_file_state(path: Path, expected: str | None) -> None:
    if file_state(path) != expected:
        raise ArchiveError(f"Monthly publication artifact changed since validation: {path}")


def _read_previous_close(path: Path, *, store: str, period: str) -> dict:
    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict) or payload.get("store") != store or payload.get("period") != period:
            raise ValueError("store or period mismatch")
        return payload
    except (ValueError, OSError) as exc:
        raise ArchiveError(f"Existing close manifest is unreadable or has conflicting identity: {path}: {exc}") from exc


def _previous_archive_records(manifest: Mapping, *, archive_root: Path) -> tuple[ArchiveRecord, ...]:
    """Return preserved evidence that exists, never a new delivery's copy plan."""
    root = Path(archive_root).resolve()
    rows = manifest.get("sources")
    if not isinstance(rows, list) or not rows:
        raise ArchiveError("Existing close manifest contains no preserved source records.")
    records = []
    for row in rows:
        try:
            archive_path = (root / row["archive_path"]).resolve()
            size = row["size_bytes"]
            if not _is_within(archive_path, root) or not archive_path.is_file():
                raise ValueError("source is missing or outside its archive")
            if isinstance(size, bool) or not isinstance(size, int) or size < 0:
                raise ValueError("invalid recorded size")
            if archive_path.stat().st_size != size or sha256_file(archive_path) != row["sha256"]:
                raise ValueError("source bytes no longer match the manifest")
            records.append(ArchiveRecord(
                role=row["role"], source_path=Path(row.get("source_path") or archive_path),
                archive_path=archive_path, archive_category=row["archive_category"],
                sha256=row["sha256"], size_bytes=size, remove_after_publish=False,
            ))
        except (KeyError, TypeError, ValueError, OSError) as exc:
            raise ArchiveError(f"Preserved monthly evidence no longer verifies: {row!r}: {exc}") from exc
    return tuple(records)


def _is_review_explanation_input(path: Path, data: _CloseData) -> bool:
    if not any(record.role == "monthly_variance_explanation" and record.source_path.resolve() == path.resolve()
               for record in data.archive_records):
        return False
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        return ("Monthly Close Report" in workbook.sheetnames
                and workbook["Monthly Close Report"]["A4"].value == CloseStatus.REVIEW_REQUIRED.value)
    finally:
        workbook.close()


def _legacy_monthly_identity(*, config: StoreConfig, fiscal_period: FiscalPeriod,
                             archive_root: Path, micros_work_dir: Path) -> str | None:
    """Recompute old manifest evidence before treating its delivery as unchanged."""
    from gift_card_recon.monthly_close_cli import _resolve_archive_reissue
    try:
        package = _resolve_archive_reissue(archive_root=archive_root, store=config.store, period=fiscal_period.period_key)
        data = _build_close_data(
            config=config, fiscal_period=fiscal_period, input_dir=package.input_dir,
            micros_path=package.micros_path, micros_work_dir=micros_work_dir,
            archive_root=archive_root, darden_path=package.job.darden_path,
            darden_report=None, cleanup_sources=False, allow_unconfigured_micros=True,
            archived_variance_explanations=dict(package.variance_explanations),
            archived_monthly_variance_explanation=package.monthly_variance_explanation,
            archived_source_paths=dict(package.source_paths),
        )
        return data.input_set_id
    except (ParseError, ArchiveError, OSError, ValueError, RuntimeError):
        # A legacy package lacking independently verifiable identity cannot be
        # called a duplicate. Replacement remains deliberate and preserves it.
        return None


def _verify_selected_archive_sources(sources: Mapping[str, Sequence[Path]], *, archive_root: Path,
                                     input_dir: Path, config: StoreConfig, fiscal_period: FiscalPeriod) -> None:
    manifest = _read_previous_close(input_dir / "close_manifest.json", store=config.store,
                                    period=fiscal_period.period_key)
    recorded: dict[str, list[Path]] = {}
    root = Path(archive_root).resolve()
    for row in manifest.get("sources", []):
        path = (root / row["archive_path"]).resolve()
        if not _is_within(path, input_dir) or not path.is_file():
            raise ArchiveError(f"Archived source is outside its package or missing: {path}")
        if path.stat().st_size != row["size_bytes"] or sha256_file(path) != row["sha256"]:
            raise ArchiveError(f"Archived source no longer matches its selected manifest: {path}")
        recorded.setdefault(row["role"], []).append(path)
    selected = {role: sorted(Path(path).resolve() for path in paths) for role, paths in sources.items()}
    if selected != {role: sorted(paths) for role, paths in recorded.items()}:
        raise ArchiveError("Selected archive sources do not exactly match the current close manifest.")
    for role, count in {"Gift Card Summary": 1, "Weekly Gift Card Activity": len(fiscal_period.expected_week_endings),
                        "Micros Daily System Totals": 1, "Micros Tender Detail": 1, "Darden Credit Memo": 1}.items():
        if len(selected.get(role, ())) != count:
            raise ArchiveError(f"Archived {role} requires {count} selected source(s).")


def _verify_existing_workbook(path: Path, manifest: Mapping) -> None:
    if not path.is_file():
        raise ArchiveError(f"Published monthly workbook is missing: {path}")
    expected = manifest.get("workbook_control_identity") or manifest.get("workbook_values_identity")
    if expected:
        actual = (_monthly_report_control_identity(path) if manifest.get("workbook_control_identity")
                  else workbook_values_identity(path))
        if actual != expected:
            raise ArchiveError("Published monthly workbook values or formulas changed; restore the retained revision before continuing.")
    else:
        # Old manifests still protect their first migration; new revisions use
        # business cells, allowing formatting and Excel metadata changes.
        artifact = next((row for row in manifest.get("artifacts", []) if row.get("role") == "workbook"), None)
        if artifact is None or sha256_file(path) != artifact.get("sha256"):
            raise ArchiveError("Legacy monthly workbook does not match its recorded artifact; verify it before migration.")


def resolve_monthly_workbook(manifest: Mapping, expected_path: Path) -> Path:
    """Resolve a local report by protected content, allowing a harmless rename."""
    expected_path = Path(expected_path)
    if expected_path.is_file():
        _verify_existing_workbook(expected_path, manifest)
        return expected_path
    matches = []
    for candidate in sorted(expected_path.parent.glob("*.xlsx")):
        if candidate.name.startswith(("~$", ".")):
            continue
        try:
            _verify_existing_workbook(candidate, manifest)
        except (ArchiveError, OSError, ValueError, BadZipFile):
            continue
        matches.append(candidate)
    if len(matches) != 1:
        raise ArchiveError(f"Expected one preserved monthly workbook in {expected_path.parent}; found {len(matches)}.")
    return matches[0]


def _resolve_monthly_pdf(manifest: Mapping, expected_path: Path) -> Path:
    record = next((row for row in manifest.get("artifacts", []) if row.get("role") == "pdf"), None)
    if record is None:
        return expected_path
    candidates = ([expected_path] if expected_path.is_file() else []) + list(expected_path.parent.glob("*.pdf"))
    matches = {path for path in candidates if path.is_file() and path.stat().st_size == record.get("size_bytes")
               and sha256_file(path) == record.get("sha256")}
    if expected_path in matches:
        return expected_path
    if len(matches) == 1:
        return matches.pop()
    return expected_path


def _portable_monthly_workbook(manifest: Mapping, *, archive_root: Path, config: StoreConfig,
                               fiscal_period: FiscalPeriod, output_root: Path | None) -> Path:
    operations = Path(archive_root).resolve().parent
    if output_root is not None:
        expected, _ = canonical_output_paths(Path(output_root), config=config, fiscal_period=fiscal_period)
    elif manifest.get("workbook_relative_path"):
        expected = (operations / manifest["workbook_relative_path"]).resolve()
        if not _is_within(expected, operations):
            raise ArchiveError("Recorded workbook path escapes this operations workspace.")
    else:
        folder = operations / ("03 Finished Reports" if (operations / "03 Finished Reports").exists() else "Output")
        expected, _ = canonical_output_paths(folder, config=config, fiscal_period=fiscal_period)
        if not expected.parent.exists():
            row = next((row for row in manifest.get("artifacts", []) if row.get("role") == "workbook"), None)
            if row is None:
                raise ArchiveError("No workbook is recorded in the close manifest.")
            expected = Path(row["path"])
    return resolve_monthly_workbook(manifest, expected)


def _monthly_report_control_identity(path: Path) -> str:
    from openpyxl import load_workbook
    from gift_card_recon.monthly_explanations import IDENTITY_SHEET, SHEET_NAME
    ignored = []
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if IDENTITY_SHEET in workbook.sheetnames:
            for row in workbook[IDENTITY_SHEET].iter_rows(min_row=6, values_only=True):
                if len(row) >= 4 and isinstance(row[3], int):
                    ignored.append(f"F{row[3]}")
    finally:
        workbook.close()
    return workbook_values_identity(path, {SHEET_NAME: ignored})


def _monthly_business_facts(data: _CloseData) -> dict[str, str]:
    facts = {f"{line.metric} activity": str(line.activity_value) for line in data.reconciliation.lines}
    facts.update({f"{line.metric} POS": str(line.pos_value) for line in data.reconciliation.lines})
    facts.update({f"{line.metric} Summary": str(line.summary_value) for line in data.reconciliation.lines})
    facts["status"] = data.assessment.status.value
    return facts


def _monthly_changed_facts(previous: Mapping, data: _CloseData) -> dict:
    old = previous.get("business_facts", {})
    current = _monthly_business_facts(data)
    changes = {key: {"previous": old.get(key), "current": value}
               for key, value in current.items() if old.get(key) != value}
    if not changes:
        changes["records"] = "Underlying transactions, control evidence or bound explanations changed; totals are unchanged."
    return changes


def retry_monthly_pdf(
    *, archive_root: Path, store: str | int, period: str,
    pdf_exporter: PdfExporter | None = None,
    output_root: Path | None = None,
) -> Path:
    """Export the committed workbook without recalculating or touching inputs."""
    config = get_store_config(store)
    fiscal = fiscal_period_for_label(period)
    manifest_path = Path(archive_root) / "Monthly Close" / config.store / fiscal.folder_name / "close_manifest.json"
    support = support_root_for(Path(archive_root))
    with writer_lock(support):
        recover_pending(support)
        manifest_state = file_state(manifest_path)
        manifest = _read_previous_close(manifest_path, store=config.store, period=fiscal.period_key)
        artifact = next((row for row in manifest.get("artifacts", []) if row.get("role") == "workbook"), None)
        if artifact is None:
            raise ArchiveError("No committed workbook is available for PDF retry.")
        workbook = _portable_monthly_workbook(manifest, archive_root=archive_root, config=config,
                                              fiscal_period=fiscal, output_root=output_root)
        workbook_state = file_state(workbook)
        _verify_existing_workbook(workbook, manifest)
        pdf = workbook.with_suffix(".pdf")
        pdf_state = file_state(pdf)
        revision = (manifest_path.parent / "revisions" / manifest["result_id"][:16]
                    if manifest.get("result_id") else None)
        revision_states = ({name: file_state(revision / name) for name in ("report.pdf", "close_manifest.json")}
                           if revision else {})
        with tempfile.TemporaryDirectory(prefix="monthly-pdf-", dir=support) as directory:
            staged = Path(directory) / pdf.name
            (pdf_exporter or export_monthly_close_report_pdf)(
                workbook_path=workbook, pdf_path=staged, expected_location_label=config.report_heading,
            )
            if not staged.is_file() or staged.stat().st_size == 0:
                raise PdfExportError("PDF export produced no usable file.")
            _require_file_state(workbook, workbook_state)
            _verify_existing_workbook(workbook, manifest)
            pdf_artifact = _attest_staged_artifact(staged, published_path=pdf)
            manifest["pdf_status"] = "complete"
            manifest["pdf_error"] = None
            manifest["artifacts"] = [row for row in manifest["artifacts"] if row.get("role") != "pdf"] + [{
                "role": "pdf", "path": str(pdf), "sha256": pdf_artifact.sha256,
                "size_bytes": pdf_artifact.size_bytes,
            }]
            staged_manifest = Path(directory) / "close_manifest.json"
            staged_manifest.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
            with PublicationTransaction(support) as transaction:
                if revision is not None:
                    retained_pdf = Path(directory) / "retained.pdf"
                    retained_manifest = Path(directory) / "retained-manifest.json"
                    shutil.copy2(staged, retained_pdf)
                    shutil.copy2(staged_manifest, retained_manifest)
                    transaction.replace(retained_pdf, revision / "report.pdf", expected_before=revision_states["report.pdf"])
                    transaction.replace(retained_manifest, revision / "close_manifest.json", expected_before=revision_states["close_manifest.json"])
                transaction.replace(staged, pdf, expected_before=pdf_state)
                transaction.replace(staged_manifest, manifest_path, expected_before=manifest_state)
                _require_file_state(workbook, workbook_state)
                transaction.commit()
        return pdf


def restore_monthly_revision(
    *, archive_root: Path, store: str | int, period: str, revision_id: str,
    approve_correction: str | None = None,
    output_root: Path | None = None,
) -> Path:
    """Select a preserved result with one approval bound to the current result."""
    if not re.fullmatch(r"[0-9a-f]{16,64}", revision_id):
        raise ValueError("A revision must be its recorded hexadecimal result identity.")
    config = get_store_config(store)
    fiscal = fiscal_period_for_label(period)
    root = Path(archive_root).resolve()
    manifest_path = root / "Monthly Close" / config.store / fiscal.folder_name / "close_manifest.json"
    support = support_root_for(root)
    with writer_lock(support):
        recover_pending(support)
        manifest_state = file_state(manifest_path)
        current = _read_previous_close(manifest_path, store=config.store, period=fiscal.period_key)
        selected_dir = manifest_path.parent / "revisions" / revision_id[:16]
        selected_manifest_path = selected_dir / "close_manifest.json"
        selected_manifest_state = file_state(selected_manifest_path)
        selected = _read_previous_close(selected_manifest_path, store=config.store, period=fiscal.period_key)
        selected_id = str(selected.get("result_id", ""))
        if not selected_id.startswith(revision_id) or not selected.get("input_set_id"):
            raise ArchiveError("Selected revision lacks a matching recorded business identity.")
        current_id = str(current.get("result_id") or sha256_file(manifest_path))
        current_artifact = next((row for row in current.get("artifacts", []) if row.get("role") == "workbook"), None)
        if current_artifact is None:
            raise ArchiveError("Current close has no recorded workbook.")
        workbook = _portable_monthly_workbook(current, archive_root=root, config=config,
                                              fiscal_period=fiscal, output_root=output_root)
        workbook_state = file_state(workbook)
        _verify_existing_workbook(workbook, current)
        pdf = workbook.with_suffix(".pdf")
        output_states = {workbook: workbook_state, pdf: file_state(pdf)}
        if current_id == selected_id:
            return workbook
        token = correction_token(current_id, selected["input_set_id"])
        if approve_correction != token:
            raise CorrectionRequired(current_id, selected["input_set_id"], {
                "restore_revision": selected_id, "previous": current.get("business_facts", {}),
                "selected": selected.get("business_facts", {}),
            })
        # Each preserved version names its own immutable original evidence.
        for source in selected.get("sources", []):
            path = (root / source["archive_path"]).resolve()
            if not _is_within(path, root) or not path.is_file():
                raise ArchiveError(f"Selected revision evidence is missing or outside its archive: {path}")
            if path.stat().st_size != source["size_bytes"] or sha256_file(path) != source["sha256"]:
                raise ArchiveError(f"Selected revision evidence no longer verifies: {path}")
        selected_workbook = selected_dir / "workbook.xlsx"
        selected_workbook_state = file_state(selected_workbook)
        _verify_existing_workbook(selected_workbook, selected)
        if workbook_values_identity(selected_workbook) != selected.get("workbook_values_identity"):
            raise ArchiveError("Selected revision workbook no longer matches its preserved values.")
        selected_pdf = selected_dir / "report.pdf"
        selected_pdf_state = file_state(selected_pdf)
        selected_states = {selected_workbook: selected_workbook_state, selected_pdf: selected_pdf_state}
        pdf_record = next((row for row in selected.get("artifacts", []) if row.get("role") == "pdf"), None)
        pdf_available = selected_pdf.is_file() and pdf_record is not None
        if pdf_available and (selected_pdf.stat().st_size != pdf_record["size_bytes"] or sha256_file(selected_pdf) != pdf_record["sha256"]):
            raise ArchiveError("Selected revision PDF no longer verifies.")
        result_id = hashlib.sha256(f"restore:{current_id}:{selected_id}".encode()).hexdigest()
        restored = dict(selected)
        restored.update({"result_id": result_id, "previous_result_id": current_id,
                         "restored_from_result_id": selected_id,
                         "generated_at": datetime.now().astimezone().isoformat(),
                         "pdf_status": "complete" if pdf_available else "pending",
                         "pdf_error": None if pdf_available else "Selected revision has no exported PDF."})
        restored["artifacts"] = []
        new_dir = manifest_path.parent / "revisions" / result_id[:16]
        with tempfile.TemporaryDirectory(prefix="monthly-restore-", dir=support) as directory:
            temporary = Path(directory)
            with PublicationTransaction(support) as transaction:
                for role, original, target, retained_name in (
                    ("workbook", selected_workbook, workbook, "workbook.xlsx"),
                    ("pdf", selected_pdf, pdf, "report.pdf"),
                ):
                    if role == "pdf" and not pdf_available:
                        if output_states[target] is not None or target.exists():
                            transaction.remove(target, expected_before=output_states[target])
                        continue
                    staged = temporary / retained_name
                    shutil.copy2(original, staged)
                    retained = temporary / ("retained-" + retained_name)
                    shutil.copy2(original, retained)
                    _require_file_state(original, selected_states[original])
                    _require_file_state(staged, selected_states[original])
                    _require_file_state(retained, selected_states[original])
                    artifact = _attest_staged_artifact(staged, published_path=target)
                    restored["artifacts"].append({"role": role, "path": str(target), "sha256": artifact.sha256, "size_bytes": artifact.size_bytes})
                    transaction.replace(retained, new_dir / retained_name, expected_before=None)
                    transaction.replace(staged, target, expected_before=output_states[target])
                staged_manifest = temporary / "close_manifest.json"
                staged_manifest.write_text(json.dumps(restored, indent=2) + "\n", encoding="utf-8")
                retained_manifest = temporary / "retained-manifest.json"
                shutil.copy2(staged_manifest, retained_manifest)
                transaction.replace(retained_manifest, new_dir / "close_manifest.json", expected_before=None)
                transaction.replace(staged_manifest, manifest_path, expected_before=manifest_state)
                _require_file_state(selected_manifest_path, selected_manifest_state)
                transaction.commit()
        return workbook


def prepare_monthly_close_readiness(
    *,
    store: str | int,
    store_config: StoreConfig | None = None,
    period: str,
    input_dir: Path,
    micros_path: Path | None,
    micros_work_dir: Path,
    archive_root: Path,
    darden_path: Path | None = None,
    fiscal_period: FiscalPeriod | None = None,
    allow_unconfigured_micros: bool = False,
    archived_variance_explanations: Mapping[
        date, ArchivedVarianceExplanationSource
    ] | None = None,
    archived_monthly_variance_explanation: ArchivedVarianceExplanationSource | None = None,
    archived_source_paths: Mapping[str, Sequence[Path]] | None = None,
    live_source_paths: Mapping[str, Sequence[Path]] | None = None,
    output_root: Path | None = None,
) -> _CloseData:
    """Read current evidence without staging, extracting, publishing or cleanup."""

    config = _resolve_monthly_store_config(store, store_config)
    canonical = fiscal_period_for_label(period)
    if fiscal_period is not None and (
        fiscal_period.period_key != canonical.period_key
        or fiscal_period.folder_name != canonical.folder_name
        or fiscal_period.start_date != canonical.start_date
        or fiscal_period.end_date != canonical.end_date
    ):
        raise ParseError(f"Fiscal period {period!r} does not match the supplied period object.")
    chosen_micros = Path(micros_path) if micros_path is not None else config.micros_default_path
    if chosen_micros.is_file():
        raise ParseError("Readiness requires an extracted Micros folder; archive extraction is an explicit preparation step.")
    data = _build_close_data(
        config=config,
        fiscal_period=canonical,
        input_dir=Path(input_dir),
        micros_path=Path(micros_path) if micros_path is not None else config.micros_default_path,
        micros_work_dir=Path(micros_work_dir),
        archive_root=Path(archive_root),
        darden_path=Path(darden_path) if darden_path is not None else None,
        darden_report=None,
        cleanup_sources=False,
        allow_unconfigured_micros=allow_unconfigured_micros,
        archived_variance_explanations=archived_variance_explanations,
        archived_monthly_variance_explanation=archived_monthly_variance_explanation,
        archived_source_paths=archived_source_paths,
        live_source_paths=live_source_paths,
        output_root=output_root,
    )
    return data


def assess_monthly_close_inputs(**kwargs) -> CloseAssessment:
    """Compatibility assessment API; all work is read-only."""
    return prepare_monthly_close_readiness(**kwargs).assessment


def write_monthly_close_diagnostic(
    *,
    store: str | int,
    period: str,
    output_root: Path,
    message: str,
    generated_at: datetime | None = None,
    pdf_exporter: PdfExporter | None = None,
) -> ReviewDiagnosticResult:
    """Create a red diagnostic for a discovery failure with known identity."""

    config = get_store_config(store)
    fiscal_period = fiscal_period_for_label(period)
    assessment = _failure_assessment(
        config.store,
        "darden_discovery",
        "Darden report discovery",
        message,
    )
    timestamp = generated_at.astimezone() if generated_at else datetime.now().astimezone()
    return _write_review_diagnostic(
        output_root=Path(output_root),
        config=config,
        fiscal_period=fiscal_period,
        assessment=assessment,
        generated_at=timestamp,
        message=message,
        pdf_exporter=pdf_exporter or export_monthly_close_report_pdf,
    )


@dataclass(frozen=True)
class _CloseData:
    certification: object
    reconciliation: ReconciliationResult
    weekly_variances: tuple[WeeklyPosVariance, ...]
    weekly_tender: Mapping[str, Decimal]
    period_tender: Decimal
    assessment: CloseAssessment
    archive_records: tuple[ArchiveRecord, ...]
    micros_evidence: MicrosEvidence
    weekly_explanations: Mapping[date, WeeklyVarianceExplanation]
    explanation_paths: tuple[Path, ...]
    input_set_id: str


def _build_close_data(
    *,
    config: StoreConfig,
    fiscal_period: FiscalPeriod,
    input_dir: Path,
    micros_path: Path,
    micros_work_dir: Path,
    archive_root: Path,
    darden_path: Path | None,
    darden_report: DardenCreditMemo | None,
    cleanup_sources: bool,
    allow_unconfigured_micros: bool,
    archived_variance_explanations: Mapping[
        date, ArchivedVarianceExplanationSource
    ] | None,
    archived_monthly_variance_explanation: ArchivedVarianceExplanationSource | None = None,
    archived_source_paths: Mapping[str, Sequence[Path]] | None = None,
    live_source_paths: Mapping[str, Sequence[Path]] | None = None,
    output_root: Path | None = None,
    canonical_xlsx: Path | None = None,
) -> _CloseData:
    if archived_source_paths:
        _verify_selected_archive_sources(archived_source_paths, archive_root=archive_root,
                                        input_dir=input_dir, config=config, fiscal_period=fiscal_period)
        summary_path = Path(archived_source_paths["Gift Card Summary"][0])
        activity_paths = list(archived_source_paths["Weekly Gift Card Activity"])
        config = replace(config,
                         micros_system_totals_file=Path(archived_source_paths["Micros Daily System Totals"][0]).name,
                         micros_tender_detail_file=Path(archived_source_paths["Micros Tender Detail"][0]).name)
    elif live_source_paths is not None:
        summaries = list(live_source_paths.get("Gift Card Summary", ()))
        activity_paths = list(live_source_paths.get("Weekly Gift Card Activity", ()))
        if len(summaries) != 1:
            raise ParseError(f"Monthly close requires exactly one Summary; found {len(summaries)}.")
        summary_path = Path(summaries[0])
    else:
        summary_path, activity_paths, _ = discover_input_files(input_dir, mode="monthly")
    if summary_path is None:
        raise ParseError("Monthly close requires exactly one Gift Card Summary workbook.")
    validate_micros_source(
        micros_path,
        config,
        allow_unconfigured_source=allow_unconfigured_micros,
    )
    resolved_micros_path = resolve_micros_export_dir(micros_path, micros_work_dir)
    if darden_report is None:
        chosen_darden = darden_path or _one_darden_path(input_dir)
    else:
        chosen_darden = darden_report.source_file
        if darden_path is not None and Path(darden_path).resolve() != Path(chosen_darden).resolve():
            raise ParseError("The supplied Darden report and Darden path identify different files.")
    initial_paths = [
        summary_path,
        *activity_paths,
        Path(chosen_darden),
        _required_source_file(resolved_micros_path, config.micros_system_totals_file),
        _required_source_file(resolved_micros_path, config.micros_tender_detail_file),
    ]
    initial_hashes = {
        path.resolve(): sha256_file(path)
        for path in initial_paths
    }

    # A verified manifest records intake period even if the source was renamed.
    if not archived_source_paths or summary_report_end(summary_path) is not None:
        _validate_summary_period(summary_path, fiscal_period)

    summary = parse_summary(summary_path, store=config.store)
    activities = [
        parse_activity_file(path, summary.conversion_promo_codes) for path in activity_paths
    ]
    activity_evidence = validate_activity_evidence(
        activities,
        store=config.store,
        period_start=fiscal_period.start_date,
        period_end=fiscal_period.end_date,
        expected_week_endings=fiscal_period.expected_week_endings,
    )

    if darden_report is None:
        darden_report = parse_darden_credit_memo(chosen_darden)
    certification = build_monthly_close_certification(
        store=config.store,
        period=fiscal_period.period_key,
        period_start=fiscal_period.start_date,
        period_end=fiscal_period.end_date,
        summary=summary,
        darden_credit_memo=darden_report,
    )

    micros_evidence = load_micros_evidence(
        resolved_micros_path,
        config=config,
        activity_evidence=activity_evidence,
        period_start=fiscal_period.start_date,
        period_end=fiscal_period.end_date,
        validate_source=False,
        allow_unconfigured_source=allow_unconfigured_micros,
    )
    weekly_variances = tuple(
        build_weekly_pos_variances(
            activity_evidence,
            micros_evidence,
            conversion_promo_codes=summary.conversion_promo_codes,
        )
    )
    pos_controls = PosControls(
        store=config.store,
        period=fiscal_period.period_key,
        pos_gift_card_issue=money(
            sum((row.pos_issue for row in weekly_variances), Decimal("0.00"))
        ),
        pos_gift_card_payment=money(
            sum((row.pos_payment for row in weekly_variances), Decimal("0.00"))
        ),
    )
    reconciliation = build_reconciliation(
        store=config.store,
        period=fiscal_period.period_key,
        period_end=fiscal_period.end_date,
        summary=summary,
        activities=list(activity_evidence.files),
        pos_controls=pos_controls,
        mode="monthly",
        additional_source_files=(
            darden_report.source_file,
            micros_evidence.system_totals_path,
            micros_evidence.tender_detail_path,
        ),
        strict_nonzero_review=True,
    )

    weekly_tender = weekly_tender_variances(
        micros_evidence,
        week_endings=fiscal_period.expected_week_endings,
    )
    period_tender = period_tender_variance(micros_evidence)
    weekly_explanations, explanation_paths, explanation_hashes, explanation_blockers = _load_weekly_variance_explanations(
        input_dir=input_dir,
        config=config,
        weekly_variances=weekly_variances,
        weekly_tender=weekly_tender,
        archived_variance_explanations=archived_variance_explanations,
        archive_root=archive_root,
        output_root=output_root,
    )
    initial_hashes.update(explanation_hashes)
    evidence_items = _evidence_items(
        config=config,
        fiscal_period=fiscal_period,
        summary_path=summary_path,
        activity_paths=activity_paths,
        darden_path=darden_report.source_file,
        micros_evidence=micros_evidence,
        archive_root=archive_root,
        cleanup_sources=cleanup_sources,
        explanation_paths=explanation_paths,
    )
    archive_records = tuple(plan_evidence_archive(evidence_items, archive_root=archive_root))
    for record in archive_records:
        initial_hash = initial_hashes.get(record.source_path.resolve())
        if initial_hash is None or initial_hash != record.sha256:
            raise ArchiveError(
                f"Evidence changed while the close was being calculated: {record.source_path}. "
                "No report was published; rerun with a stable source file."
            )

    integrity_controls = (
        integrity_control(
            code="summary_identity",
            label="Summary identity and required values",
            passed=True,
            pass_message="Exactly one Summary row matched this store and required money parsed.",
            failure_message="Summary identity validation failed.",
        ),
        integrity_control(
            code="activity_identity",
            label="Activity report identity",
            passed=True,
            pass_message="Every activity report identifies the configured store.",
            failure_message="Activity identity validation failed.",
        ),
        integrity_control(
            code="activity_coverage",
            label="Activity weekly coverage",
            passed=True,
            pass_message="Exactly one non-overlapping Monday-Sunday report covers every expected week.",
            failure_message="Activity coverage validation failed.",
        ),
        integrity_control(
            code="darden_identity",
            label="Darden identity and period",
            passed=True,
            pass_message="The Darden memo store and service dates match this close.",
            failure_message="Darden identity validation failed.",
        ),
        integrity_control(
            code="micros_source",
            label="Micros source and layout",
            passed=True,
            pass_message=f"Validated {config.micros_source_label} and configured control columns.",
            failure_message="Micros source validation failed.",
        ),
        integrity_control(
            code="micros_coverage",
            label="Micros date coverage",
            passed=True,
            pass_message="Every fiscal date is present or is an evidence-confirmed scheduled closure.",
            failure_message="Micros coverage validation failed.",
        ),
        integrity_control(
            code="tender_evidence",
            label="Tender evidence",
            passed=True,
            pass_message="Tender Detail is present, well formed, normalized, and assessed.",
            failure_message="Tender evidence validation failed.",
        ),
        integrity_control(
            code="archive_integrity",
            label="Archive plan and source hashes",
            passed=True,
            pass_message="Every required source was hashed and assigned a canonical archive path.",
            failure_message="Archive preflight failed.",
        ),
    )
    summary_activity = {
        line.metric: line.activity_variance for line in reconciliation.lines
    }
    weekly_pos: dict[str, Decimal] = {}
    for row in weekly_variances:
        week = row.week_ending.strftime("%m/%d/%Y") if row.week_ending else "Unknown"
        weekly_pos[f"Week ending {week} POS issue"] = row.issue_variance
        weekly_pos[f"Week ending {week} POS payment"] = row.payment_variance
        weekly_pos[f"Week ending {week} POS net"] = row.net_variance
    period_pos = {
        f"Period POS {line.metric}": line.pos_variance for line in reconciliation.lines
    }
    assessment = assess_monthly_close(
        store=config.store,
        darden_variance=certification.variance,
        summary_activity_variances=summary_activity,
        weekly_pos_variances=weekly_pos,
        period_pos_variances=period_pos,
        weekly_tender_variances=weekly_tender,
        period_tender_variances={"Period tender": period_tender},
        integrity_controls=(*integrity_controls, *explanation_blockers),
        additional_required_integrity_codes=("archive_integrity",),
        expected_week_count=len(fiscal_period.expected_week_endings),
    )
    explained_variances = _bound_weekly_explanations(
        assessment, weekly_variances, weekly_explanations
    )
    monthly_source = None
    if archived_monthly_variance_explanation is not None:
        monthly_source = archived_monthly_variance_explanation.path
        if (
            not _is_within(monthly_source, input_dir)
            or monthly_source.parent.name.casefold() != "monthly variance explanations"
            or not monthly_source.is_file()
        ):
            raise ParseError("Archived monthly variance explanation is outside its verified close package or missing.")
        if monthly_source.stat().st_size != archived_monthly_variance_explanation.size_bytes:
            raise ParseError("Archived monthly variance explanation size no longer matches its close manifest.")
        if sha256_file(monthly_source) != archived_monthly_variance_explanation.sha256:
            raise ParseError("Archived monthly variance explanation hash no longer matches its close manifest.")
    elif archived_variance_explanations is None and any(
        control.variance not in (None, Decimal("0.00"))
        and (control.code == "darden_summary_match" or control.code.startswith(("summary_activity_", "period_pos_", "period_tender_")))
        for control in assessment.controls
    ):
        monthly_source = locate_monthly_variance_explanation_source(
            input_dir=input_dir, store=config.store, period=fiscal_period.period_key,
            output_root=output_root, canonical_path=canonical_xlsx,
        )
    if monthly_source is not None:
        source_hash = sha256_file(monthly_source)
        try:
            monthly_explanations = read_monthly_variance_explanations(
                monthly_source, store=config.store, period=fiscal_period.period_key,
                controls=assessment.controls,
            )
        except BadZipFile as exc:
            if archived_monthly_variance_explanation is not None:
                raise ParseError("Archived monthly variance explanation workbook is unreadable.") from exc
            monthly_explanations = {}
            integrity_controls += (ControlOutcome(
                code="monthly_explanation_evidence", label="Monthly explanation source",
                disposition=ControlDisposition.BLOCK,
                message=f"The existing monthly explanation workbook is unreadable: {monthly_source}.",
            ),)
        if source_hash != sha256_file(monthly_source):
            raise ArchiveError("Monthly variance explanations changed while they were being read; save, close, and rerun.")
        if monthly_explanations:
            explained_variances.update(monthly_explanations)
            initial_hashes[monthly_source.resolve()] = source_hash
            evidence_items.append(EvidenceItem(
                "monthly_variance_explanation", monthly_source,
                str(Path("Monthly Close") / config.store / fiscal_period.folder_name / "Monthly Variance Explanations"),
                False,
            ))
    if explained_variances or any(control.code == "monthly_explanation_evidence" for control in integrity_controls):
        assessment = assess_monthly_close(
            store=config.store, darden_variance=certification.variance,
            summary_activity_variances=summary_activity, weekly_pos_variances=weekly_pos,
            period_pos_variances=period_pos, weekly_tender_variances=weekly_tender,
            period_tender_variances={"Period tender": period_tender},
            integrity_controls=(*integrity_controls, *explanation_blockers),
            additional_required_integrity_codes=("archive_integrity",),
            expected_week_count=len(fiscal_period.expected_week_endings),
            explained_variances=explained_variances,
        )
    archive_records = tuple(plan_evidence_archive(evidence_items, archive_root=archive_root))
    for record in archive_records:
        if initial_hashes.get(record.source_path.resolve()) != record.sha256:
            raise ArchiveError(f"Evidence changed while the close was being calculated: {record.source_path}.")
    assessment_exceptions = [
        (
            control.disposition.value,
            f"{control.label}: {control.message}",
        )
        for control in assessment.controls
        if not control.passed
    ]
    reconciliation = replace(reconciliation, exceptions=assessment_exceptions)
    input_identity = input_set_identity(
        activities, summary=summary,
        extra={
            "store": config.store, "period": fiscal_period.period_key,
            "darden": {name: value for name, value in vars(darden_report).items() if name != "source_file"},
            "daily_pos": micros_evidence.daily_pos,
            "daily_tender": micros_evidence.daily_tender,
            "daily_tender_magnitude": micros_evidence.daily_tender_magnitude,
            "tender_observed_dates": micros_evidence.tender_observed_dates,
            "accepted_closed_dates": micros_evidence.accepted_closed_dates,
            "explanations": explained_variances,
        },
    )
    return _CloseData(
        certification=certification,
        reconciliation=reconciliation,
        weekly_variances=weekly_variances,
        weekly_tender=weekly_tender,
        period_tender=period_tender,
        assessment=assessment,
        archive_records=archive_records,
        micros_evidence=micros_evidence,
        weekly_explanations=weekly_explanations,
        explanation_paths=explanation_paths,
        input_set_id=input_identity,
    )


def _load_weekly_variance_explanations(
    *,
    input_dir: Path,
    config: StoreConfig,
    weekly_variances: Sequence[WeeklyPosVariance],
    weekly_tender: Mapping[str, Decimal],
    archived_variance_explanations: Mapping[
        date, ArchivedVarianceExplanationSource
    ] | None,
    archive_root: Path | None = None,
    output_root: Path | None = None,
) -> tuple[
    dict[date, WeeklyVarianceExplanation],
    tuple[Path, ...],
    dict[Path, str],
    tuple[ControlOutcome, ...],
]:
    archived_sources = (
        None
        if archived_variance_explanations is None
        else dict(archived_variance_explanations)
    )
    expected_week_endings = {
        row.week_ending
        for row in weekly_variances
        if row.week_ending is not None and row.report_begin is not None
    }
    if archived_sources is not None:
        if any(not isinstance(week_end, date) for week_end in archived_sources):
            raise ParseError("Archived variance explanation mapping contains an invalid week ending.")
        unexpected = sorted(set(archived_sources) - expected_week_endings)
        if unexpected:
            labels = ", ".join(week_end.isoformat() for week_end in unexpected)
            raise ParseError(
                "Archived variance explanation mapping contains unexpected week ending(s): "
                f"{labels}."
            )
        claimed_paths: dict[Path, date] = {}
        input_root = Path(input_dir).resolve(strict=False)
        for week_end, source in archived_sources.items():
            if not isinstance(source, ArchivedVarianceExplanationSource):
                raise ParseError(
                    "Archived variance explanation mapping contains an invalid source record."
                )
            resolved_path = source.path.resolve(strict=False)
            if (
                resolved_path.parent.parent != input_root
                or resolved_path.parent.name.casefold() != "variance explanations"
            ):
                raise ParseError(
                    "Archived weekly variance explanation is outside its canonical "
                    f"store-period folder: {source.path}"
                )
            prior_week = claimed_paths.get(resolved_path)
            if prior_week is not None:
                raise ParseError(
                    "Archived weekly variance explanation is assigned to more than one "
                    f"week: {resolved_path}"
                )
            claimed_paths[resolved_path] = week_end

    explanations: dict[date, WeeklyVarianceExplanation] = {}
    paths: list[Path] = []
    hashes: dict[Path, str] = {}
    blockers: list[ControlOutcome] = []
    for row in weekly_variances:
        if row.week_ending is None or row.report_begin is None:
            continue
        week_label = f"Week ending {row.week_ending:%m/%d/%Y}"
        tender = money(weekly_tender[f"{week_label} tender"])
        control_values = {
            "pos_issue_variance": money(row.issue_variance),
            "pos_payment_variance": money(row.payment_variance),
            "pos_net_variance": money(row.net_variance),
            "tender_variance": tender,
        }
        controls = (
            ("POS issue", control_values["pos_issue_variance"]),
            ("POS payment", control_values["pos_payment_variance"]),
            ("POS net", control_values["pos_net_variance"]),
            ("tender", control_values["tender_variance"]),
        )
        over_limit = tuple(
            (label, value)
            for label, value in controls
            if abs(value) > REVIEW_VARIANCE_LIMIT
        )
        expected_path = variance_explanation_path(input_dir, config.store, row.week_ending)
        archived_source = (
            None if archived_sources is None else archived_sources.get(row.week_ending)
        )
        path = (
            archived_source.path if archived_source is not None else
            (resolve_live_weekly_explanation_path(
                input_dir, config.store, row.week_ending,
                archive_root=archive_root, output_root=output_root,
            ) or expected_path)
            if archived_sources is None else expected_path
        )
        source_is_present = (
            path.is_file()
            if archived_sources is None or archived_source is not None
            else False
        )
        if not source_is_present:
            if archived_source is not None:
                raise ParseError(f"Archived weekly variance explanation is missing: {path}")
            if over_limit:
                details = ", ".join(
                    f"{label} {value:+,.2f}" for label, value in over_limit
                )
                archive_context = (
                    " The verified archive manifest contains no explanation form for "
                    "this week."
                    if archived_sources is not None
                    else " The weekly runner creates this form automatically."
                )
                message = (
                    f"Weekly variance explanation is required for store {config.store} "
                    f"week ending {row.week_ending:%m/%d/%Y} because {details} exceeds "
                    f"the ${REVIEW_VARIANCE_LIMIT:.2f} limit. Expected {path}." + archive_context
                )
                if archived_sources is not None:
                    raise ParseError(message)
                blockers.append(ControlOutcome(
                    code=f"weekly_explanation_{row.week_ending:%Y%m%d}",
                    label=f"Week ending {row.week_ending:%m/%d/%Y} explanation",
                    disposition=ControlDisposition.BLOCK, message=message,
                ))
            continue
        if archived_source is not None and path.stat().st_size != archived_source.size_bytes:
            raise ParseError(
                "Archived weekly variance explanation size no longer matches its close "
                f"manifest: {path}"
            )
        initial_hash = sha256_file(path)
        if archived_source is not None and initial_hash != archived_source.sha256:
            raise ParseError(
                "Archived weekly variance explanation hash no longer matches its close "
                f"manifest: {path}"
            )
        explanation = read_variance_explanation_workbook(
            path,
            expected_store=config.store,
            expected_week_start=row.report_begin,
            expected_week_end=row.week_ending,
            require_text=False,
        )
        mismatch_details = variance_control_mismatch_details(
            explanation,
            **control_values,
        )
        if mismatch_details:
            raise ParseError(
                "Weekly variance explanation controls do not match current source "
                f"evidence for store {config.store} week ending "
                f"{row.week_ending:%m/%d/%Y} at {path}: "
                + "; ".join(mismatch_details)
                + ". Resolve the evidence change before monthly close."
            )
        final_hash = sha256_file(path)
        if final_hash != initial_hash:
            raise ParseError(
                f"Weekly variance explanation changed while it was being read: {path}. "
                "Save and close the workbook, then rerun."
            )
        paths.append(path)
        hashes[path.resolve()] = initial_hash
        if explanation.explanation.strip():
            explanations[row.week_ending] = explanation
        elif over_limit:
            # Defensive guard if a future reader implementation stops enforcing
            # require_text itself.
            message = (
                f"Enter a plain-text weekly variance explanation in "
                f"{EXPLANATION_SHEET}!{EXPLANATION_INPUT_CELL} at {path}, save, and rerun."
            )
            if archived_sources is not None:
                raise ParseError(message)
            blockers.append(ControlOutcome(
                code=f"weekly_explanation_{row.week_ending:%Y%m%d}",
                label=f"Week ending {row.week_ending:%m/%d/%Y} explanation",
                disposition=ControlDisposition.BLOCK, message=message,
            ))
    return explanations, tuple(paths), hashes, tuple(blockers)


def _bound_weekly_explanations(
    assessment: CloseAssessment,
    weekly_variances: Sequence[WeeklyPosVariance],
    weekly_explanations: Mapping[date, WeeklyVarianceExplanation],
) -> dict[str, str]:
    """Authorize only controls covered by current, identity-checked weekly evidence."""
    by_label = {control.label: control for control in assessment.controls}
    explained: dict[str, str] = {}
    dimensions = {
        "POS issue": "Period POS Gift Card Issue / Activations",
        "POS payment": "Period POS Gift Card Payment / Redemptions",
        "POS net": "Period POS Net Gift Card Impact",
        "tender": "Period tender",
    }
    for dimension, period_label in dimensions.items():
        contributions: list[ControlOutcome] = []
        covered_weeks: list[str] = []
        for row in weekly_variances:
            if row.week_ending is None:
                continue
            label = f"Week ending {row.week_ending:%m/%d/%Y} {dimension}"
            control = by_label[label]
            if control.variance is None or control.variance == 0:
                continue
            contributions.append(control)
            narrative = weekly_explanations.get(row.week_ending)
            if narrative is not None and narrative.explanation.strip():
                explained[control.code] = narrative.explanation
                covered_weeks.append(f"{row.week_ending:%m/%d/%Y}")
        period = by_label.get(period_label)
        if (
            period is not None and period.variance not in (None, Decimal("0.00"))
            and contributions
            and all(control.code in explained for control in contributions)
            and sum((control.variance for control in contributions), Decimal("0.00")) == period.variance
        ):
            explained[period.code] = (
                "Covered by the matching weekly explanations retained for "
                + ", ".join(covered_weeks) + "."
            )
    return explained


def _evidence_items(
    *,
    config: StoreConfig,
    fiscal_period: FiscalPeriod,
    summary_path: Path,
    activity_paths: Sequence[Path],
    darden_path: Path,
    micros_evidence: MicrosEvidence,
    archive_root: Path,
    cleanup_sources: bool,
    explanation_paths: Sequence[Path] = (),
) -> list[EvidenceItem]:
    base = Path("Monthly Close") / config.store / fiscal_period.folder_name
    archive_package = Path(archive_root).resolve(strict=False) / base

    def removable(path: Path) -> bool:
        return cleanup_sources and not _is_within(path, archive_root)

    def explanation_category(path: Path) -> str:
        if _is_within(path, archive_package):
            archive_base = Path(archive_root).resolve(strict=False)
            source_parent = Path(path).resolve(strict=False).parent
            try:
                return str(source_parent.relative_to(archive_base))
            except ValueError:  # pragma: no cover - guarded by _is_within
                pass
        return str(base / "Variance Explanations")

    items = [
        EvidenceItem("Gift Card Summary", summary_path, str(base / "summary"), removable(summary_path)),
        *[
            EvidenceItem("Weekly Gift Card Activity", path, str(base / "activity"), removable(path))
            for path in activity_paths
        ],
        *[
            EvidenceItem(
                "Weekly Variance Explanation",
                path,
                explanation_category(path),
                removable(path) and path.parent.name.casefold() == "variance explanations",
            )
            for path in explanation_paths
        ],
        EvidenceItem("Darden Credit Memo", darden_path, str(base / "darden"), removable(darden_path)),
        EvidenceItem("Micros Daily System Totals", micros_evidence.system_totals_path, str(base / "micros"), False),
        EvidenceItem("Micros Tender Detail", micros_evidence.tender_detail_path, str(base / "micros"), False),
    ]
    return items


def _write_full_workbook(
    *,
    output_path: Path,
    result: ReconciliationResult,
    close_data: _CloseData,
    generated_at: datetime,
    archive_published: bool = False,
) -> None:
    assessment = close_data.assessment
    weekly_rows, weekly_codes = _weekly_report_rows(close_data)
    period_pos = next(
        (line.pos_variance for line in result.lines if line.metric == "Net Gift Card Impact"),
        None,
    )
    write_reconciliation_workbook(
        result,
        output_path,
        monthly_close_certification=close_data.certification,
        close_assessment=assessment,
        weekly_pos_variances=list(close_data.weekly_variances),
        weekly_close_rows=weekly_rows,
        period_pos_net_variance=period_pos,
        period_pos_disposition=_worst_disposition(
            control for control in assessment.controls if control.code.startswith("period_pos_")
        ),
        period_tender_variance=close_data.period_tender,
        period_tender_disposition=_worst_disposition(
            control for control in assessment.controls if control.code.startswith("period_tender_")
        ),
        evidence_notes=(
            (
                "Canonical archive-relative paths and SHA-256 hashes are recorded in Source Files and the close manifest."
                if archive_published
                else "Source hashes and archive destinations were planned, but no canonical archive or close manifest was published for this diagnostic."
            ),
            "Scheduled Mondays are accepted only when both activity and tender evidence are zero; existing Monday POS is included normally.",
        ),
        source_labels=(
            DEFAULT_EVIDENCE_LABELS
            + (("Weekly Variance Explanations",) if close_data.explanation_paths else ())
        ),
        weekly_control_codes=frozenset(weekly_codes),
        generated_at=generated_at,
        micros_source_label=assessment.store_config.micros_source_label,
    )


def _weekly_report_rows(
    close_data: _CloseData,
) -> tuple[list[WeeklyCloseReportRow], set[str]]:
    assessment = close_data.assessment
    control_by_label = {control.label: control for control in assessment.controls}
    result: list[WeeklyCloseReportRow] = []
    codes: set[str] = set()
    for row in close_data.weekly_variances:
        week = row.week_ending.strftime("%m/%d/%Y") if row.week_ending else "Unknown"
        labels = (
            f"Week ending {week} POS issue",
            f"Week ending {week} POS payment",
            f"Week ending {week} POS net",
            f"Week ending {week} tender",
        )
        controls = [control_by_label[label] for label in labels]
        explanation_control = control_by_label.get(f"Week ending {week} explanation")
        if explanation_control is not None:
            controls.append(explanation_control)
        codes.update(control.code for control in controls)
        non_pass = [
            f"{control.label.rsplit(' ', 2)[-2]} {control.label.rsplit(' ', 1)[-1]} "
            f"{control.variance:+,.2f}"
            if control.variance is not None
            else control.label
            for control in controls
            if not control.passed
        ]
        explanation = close_data.weekly_explanations.get(row.week_ending)
        result.append(
            WeeklyCloseReportRow(
                week_ending=row.week_ending,
                coverage=row.coverage_status,
                pos_issue_variance=row.issue_variance,
                pos_payment_variance=row.payment_variance,
                pos_net_variance=row.net_variance,
                tender_variance=close_data.weekly_tender[labels[-1]],
                disposition=_worst_disposition(controls),
                evidence_note=(
                    ("Explained: " if all(not control.is_blocking and not control.needs_review for control in controls)
                     else "Review required: ") + "; ".join(non_pass)
                    if non_pass
                    else "Evidence complete; all weekly controls passed."
                ),
                variance_explanation=(
                    explanation.explanation if explanation is not None else ""
                ),
            )
        )
    return result, codes


def _with_archive_source_audits(
    result: ReconciliationResult,
    records: Sequence[ArchiveRecord],
) -> ReconciliationResult:
    audits: list[SourceFileAudit] = []
    for record in records:
        relative_path = Path(record.archive_category) / record.archive_path.name
        audits.append(
            SourceFileAudit(
                path=relative_path,
                file_type=record.source_path.suffix.lower().lstrip("."),
                size_bytes=record.size_bytes,
                modified_at=file_modified_at(record.source_path),
                sha256=record.sha256,
            )
        )
    return replace(result, source_files=audits)


def _write_detailed_review(
    *,
    output_root: Path,
    config: StoreConfig,
    fiscal_period: FiscalPeriod,
    close_data: _CloseData,
    generated_at: datetime,
    pdf_exporter: PdfExporter,
) -> ReviewDiagnosticResult:
    unreadable_source = next((
        control for control in close_data.assessment.controls
        if control.code == "monthly_explanation_evidence"
    ), None)
    if unreadable_source is not None:
        # Preserve unreadable operator inputs using the retention-aware path,
        # while keeping the known monetary blockers in the full assessment.
        return _write_review_diagnostic(
            output_root=output_root, config=config, fiscal_period=fiscal_period,
            assessment=close_data.assessment, generated_at=generated_at,
            message=unreadable_source.message, pdf_exporter=pdf_exporter,
        )
    review_xlsx, review_pdf = review_output_paths(
        output_root,
        config=config,
        fiscal_period=fiscal_period,
    )
    review_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with _publication_staging_directory(
        review_xlsx.parent,
        prefix="review-required-",
        support_root=Path(output_root).parent / "_automation_runs" / "support",
    ) as temp_dir:
        temp_xlsx = Path(temp_dir) / review_xlsx.name
        temp_pdf = Path(temp_dir) / review_pdf.name
        _write_full_workbook(
            output_path=temp_xlsx,
            result=close_data.reconciliation,
            close_data=close_data,
            generated_at=generated_at,
        )
        pdf_error = _try_export_review(
            temp_xlsx,
            temp_pdf,
            config=config,
            pdf_exporter=pdf_exporter,
        )
        if pdf_error is None:
            _publish_diagnostic(
                temp_xlsx=temp_xlsx,
                temp_pdf=temp_pdf,
                canonical_xlsx=review_xlsx,
                canonical_pdf=review_pdf,
            )
            return ReviewDiagnosticResult(review_xlsx, review_pdf)
        _publish_diagnostic(
            temp_xlsx=temp_xlsx,
            canonical_xlsx=review_xlsx,
            canonical_pdf=review_pdf,
            temp_pdf=None,
        )
        return ReviewDiagnosticResult(review_xlsx, None, pdf_error)


def _write_review_diagnostic(
    *,
    output_root: Path,
    config: StoreConfig,
    fiscal_period: FiscalPeriod,
    assessment: CloseAssessment,
    generated_at: datetime,
    message: str,
    pdf_exporter: PdfExporter,
) -> ReviewDiagnosticResult:
    review_xlsx, review_pdf = review_output_paths(
        output_root,
        config=config,
        fiscal_period=fiscal_period,
    )
    review_xlsx.parent.mkdir(parents=True, exist_ok=True)
    evidence_notes = ["No inputs were archived and no canonical close report was published."]
    expected_review_hash: str | None = None
    if review_xlsx.is_file():
        from openpyxl import load_workbook
        from gift_card_recon.monthly_explanations import SHEET_NAME as monthly_input_sheet

        expected_review_hash = sha256_file(review_xlsx)
        # Preserve entered text even when malformed identity or formula input
        # caused validation to fail. This is retention, never authorization.
        try:
            prior = load_workbook(review_xlsx, data_only=False, keep_links=False)
            try:
                has_inputs = monthly_input_sheet in prior.sheetnames and any(
                    value is not None and str(value).strip()
                    for (value,) in prior[monthly_input_sheet].iter_rows(
                        min_row=7, min_col=6, max_col=6, values_only=True
                    )
                )
            finally:
                prior.close()
        except (BadZipFile, KeyError, ValueError, TypeError):
            # Unreadable prior reports may still contain operator work.
            has_inputs = True
        if sha256_file(review_xlsx) != expected_review_hash:
            raise ArchiveError("The existing monthly review changed while preserving explanation inputs; save, close, and rerun.")
        if has_inputs:
            saved_dir = review_xlsx.parent / "Saved Explanation Inputs"
            saved_dir.mkdir(parents=True, exist_ok=True)
            saved_path = saved_dir / f"inputs_{expected_review_hash[:12]}.xlsx"
            expected_size = review_xlsx.stat().st_size

            def verify_saved(path: Path) -> None:
                if not path.is_file() or path.stat().st_size != expected_size or sha256_file(path) != expected_review_hash:
                    raise ArchiveError(f"Saved monthly explanation inputs failed size/SHA-256 verification: {path}")

            if saved_path.exists():
                verify_saved(saved_path)
            else:
                handle, temporary_name = tempfile.mkstemp(prefix=".inputs-", suffix=".tmp", dir=saved_dir)
                temporary_copy = Path(temporary_name)
                try:
                    with os.fdopen(handle, "wb") as target, review_xlsx.open("rb") as source:
                        for block in iter(lambda: source.read(1024 * 1024), b""):
                            target.write(block)
                        target.flush()
                        os.fsync(target.fileno())
                    verify_saved(temporary_copy)
                    if sha256_file(review_xlsx) != expected_review_hash:
                        raise ArchiveError("The existing monthly review changed before its explanations were retained.")
                    try:
                        # A hard link atomically installs the completed file and
                        # refuses to replace an existing destination on any OS.
                        os.link(temporary_copy, saved_path)
                    except FileExistsError:
                        verify_saved(saved_path)
                    verify_saved(saved_path)
                finally:
                    temporary_copy.unlink(missing_ok=True)
            evidence_notes.append(f"Previously entered monthly explanation inputs were retained at {saved_path}. Review the saved copy before re-entering reasons for corrected amounts.")
    report_data = MonthlyCloseReportData(
        assessment=assessment,
        period=fiscal_period.period_key,
        period_start=fiscal_period.start_date,
        period_end=fiscal_period.end_date,
        generated_at=generated_at,
        explicit_exceptions=(("BLOCK", message),),
        evidence_notes=tuple(evidence_notes),
    )
    with _publication_staging_directory(
        review_xlsx.parent,
        prefix="review-required-",
        support_root=Path(output_root).parent / "_automation_runs" / "support",
    ) as temp_dir:
        temp_xlsx = Path(temp_dir) / review_xlsx.name
        temp_pdf = Path(temp_dir) / review_pdf.name
        write_monthly_close_report_workbook(report_data, temp_xlsx)
        pdf_error = _try_export_review(
            temp_xlsx,
            temp_pdf,
            config=config,
            pdf_exporter=pdf_exporter,
        )
        if expected_review_hash is not None and (
            not review_xlsx.is_file() or sha256_file(review_xlsx) != expected_review_hash
        ):
            raise ArchiveError("The monthly review changed before diagnostic publication; its existing contents were left in place.")
        if pdf_error is None:
            _publish_diagnostic(
                temp_xlsx=temp_xlsx,
                temp_pdf=temp_pdf,
                canonical_xlsx=review_xlsx,
                canonical_pdf=review_pdf,
            )
            return ReviewDiagnosticResult(review_xlsx, review_pdf)
        _publish_diagnostic(
            temp_xlsx=temp_xlsx,
            canonical_xlsx=review_xlsx,
            canonical_pdf=review_pdf,
            temp_pdf=None,
        )
        return ReviewDiagnosticResult(review_xlsx, None, pdf_error)


def _publish_diagnostic(*, temp_xlsx: Path, temp_pdf: Path | None,
                        canonical_xlsx: Path, canonical_pdf: Path) -> None:
    output_root = canonical_xlsx.parent.parent
    support = output_root.parent / "_automation_runs" / "support"
    with writer_lock(support):
        recover_pending(support)
        for target in (canonical_xlsx, canonical_pdf):
            _assert_publishable(target)
        workbook = _attest_staged_artifact(temp_xlsx, published_path=canonical_xlsx)
        pdf = _attest_staged_artifact(temp_pdf, published_path=canonical_pdf) if temp_pdf else None
        with PublicationTransaction(support) as transaction:
            transaction.replace(temp_xlsx, canonical_xlsx)
            if temp_pdf is not None:
                transaction.replace(temp_pdf, canonical_pdf)
            elif canonical_pdf.exists():
                transaction.remove(canonical_pdf)
            _verify_published_artifact(workbook)
            if pdf is not None:
                _verify_published_artifact(pdf)
            transaction.commit()


def _try_export_review(
    workbook: Path,
    pdf: Path,
    *,
    config: StoreConfig,
    pdf_exporter: PdfExporter,
) -> str | None:
    try:
        pdf_exporter(
            workbook_path=workbook,
            pdf_path=pdf,
            expected_location_label=config.report_heading,
        )
        return None
    except (PdfExportError, OSError, RuntimeError) as exc:
        return str(exc)




def _attest_staged_artifact(source: Path, *, published_path: Path) -> ManifestArtifact:
    """Bind digest and byte count to the same open staged-file instance."""

    if not source.is_file():
        raise OSError(f"Verified publication artifact is missing or empty: {source}")
    try:
        digest, size_bytes = _read_artifact_fingerprint(source)
    except OSError as exc:
        raise OSError(f"Could not attest staged publication artifact {source}: {exc}") from exc
    if size_bytes <= 0:
        raise OSError(f"Verified publication artifact is missing or empty: {source}")
    return ManifestArtifact(
        path=published_path,
        sha256=digest,
        size_bytes=size_bytes,
    )


def _verify_published_artifact(artifact: ManifestArtifact) -> None:
    path = artifact.path
    if not path.is_file():
        raise OSError(f"Published artifact is missing or empty: {path}")
    try:
        actual_hash, actual_size = _read_artifact_fingerprint(path)
    except OSError as exc:
        raise OSError(f"Could not verify published artifact {path}: {exc}") from exc
    if actual_size <= 0:
        raise OSError(f"Published artifact is missing or empty: {path}")
    if actual_size != artifact.size_bytes or actual_hash != artifact.sha256:
        raise OSError(f"Published artifact changed during publication: {path}")


def _read_artifact_fingerprint(path: Path) -> tuple[str, int]:
    """Compute digest and byte count from one open file handle."""

    digest = hashlib.sha256()
    size_bytes = 0
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
            size_bytes += len(chunk)
    return digest.hexdigest(), size_bytes




def _assert_publishable(path: Path) -> None:
    if not path.exists():
        return
    def probe() -> None:
        with path.open("r+b"):
            pass

    try:
        _retry_transient_file_operation(probe)
    except FileNotFoundError:
        return
    except OSError as exc:
        raise PermissionError(
            f"Cannot replace locked canonical output {path}. Close it and rerun; "
            "no alternate filename will be created."
        ) from exc










def _failure_assessment(store: str, code: str, label: str, message: str) -> CloseAssessment:
    return build_close_assessment(
        store=store,
        darden_variance=None,
        controls=(
            ControlOutcome(
                code=code,
                label=label,
                disposition=ControlDisposition.BLOCK,
                message=message or "Required evidence validation failed.",
            ),
        ),
    )


def _with_blocking_control(
    assessment: CloseAssessment,
    *,
    code: str,
    label: str,
    message: str,
) -> CloseAssessment:
    controls = tuple(control for control in assessment.controls if control.code != code)
    return CloseAssessment(
        store_config=assessment.store_config,
        darden_matched=assessment.darden_matched,
        controls=(
            *controls,
            ControlOutcome(
                code=code,
                label=label,
                disposition=ControlDisposition.BLOCK,
                message=message or "Publication integrity failed.",
            ),
        ),
    )


def _worst_disposition(controls: Iterable[ControlOutcome]) -> ControlDisposition:
    dispositions = {control.disposition for control in controls}
    if ControlDisposition.BLOCK in dispositions:
        return ControlDisposition.BLOCK
    if ControlDisposition.REVIEW in dispositions:
        return ControlDisposition.REVIEW
    if ControlDisposition.EXPLAINED in dispositions:
        return ControlDisposition.EXPLAINED
    return ControlDisposition.PASS


def _one_darden_path(input_dir: Path) -> Path:
    candidates = sorted((Path(input_dir) / "darden").glob("*.pdf"))
    candidates.extend(sorted(Path(input_dir).glob("*Darden*.pdf")))
    unique = list(dict.fromkeys(path.resolve() for path in candidates))
    if len(unique) != 1:
        raise ParseError(
            f"Monthly close requires exactly one Darden PDF in {input_dir}; found {len(unique)}."
        )
    return Path(unique[0])


def _validate_summary_period(path: Path, fiscal_period: FiscalPeriod) -> None:
    report_end = summary_report_end(path)
    if report_end is None:
        raise ParseError(f"Summary period is unavailable in its content or recorded intake metadata: {path.name}")
    if report_end != fiscal_period.end_date:
        raise ParseError(
            f"Summary {Path(path).name} is for report end {report_end:%Y-%m-%d}; "
            f"{fiscal_period.period_key} ends {fiscal_period.end_date:%Y-%m-%d}."
        )


def _required_source_file(folder: Path, expected_name: str) -> Path:
    matches = [
        path
        for path in Path(folder).iterdir()
        if path.is_file() and path.name.casefold() == expected_name.casefold()
    ] if Path(folder).is_dir() else []
    if len(matches) != 1:
        raise ParseError(
            f"Expected exactly one {expected_name} in {folder}; found {len(matches)}."
        )
    return matches[0]


def _is_within(path: Path, root: Path) -> bool:
    try:
        Path(path).resolve().relative_to(Path(root).resolve())
        return True
    except ValueError:
        return False




def _archive_stale_review_artifacts(
    *,
    output_root: Path,
    archive_root: Path,
    config: StoreConfig,
    fiscal_period: FiscalPeriod,
) -> list[str]:
    review_xlsx, review_pdf = review_output_paths(
        output_root,
        config=config,
        fiscal_period=fiscal_period,
    )
    legacy_review_folder = Path(output_root) / "Review Required"
    legacy_review_xlsx = legacy_review_folder / review_xlsx.name
    legacy_review_pdf = legacy_review_folder / review_pdf.name
    candidates = [
        path
        for path in (review_xlsx, review_pdf, legacy_review_xlsx, legacy_review_pdf)
        if path.is_file()
    ]
    if not candidates:
        return []
    category = str(
        Path("Generated Reports")
        / "Diagnostics"
        / config.store
        / fiscal_period.period_key
    )
    support = support_root_for(archive_root)
    try:
        records = plan_evidence_archive([
            EvidenceItem("Superseded review diagnostic", path, category, True)
            for path in candidates
        ], archive_root=archive_root)
        with writer_lock(support):
            recover_pending(support)
            with tempfile.TemporaryDirectory(prefix="review-retire-", dir=support) as directory:
                with PublicationTransaction(support) as transaction:
                    for index, record in enumerate(records):
                        if record.archive_path.exists():
                            if sha256_file(record.archive_path) != record.sha256:
                                raise ArchiveError(f"Diagnostic archive changed: {record.archive_path}")
                            continue
                        staged = Path(directory) / str(index)
                        shutil.copy2(record.source_path, staged)
                        if sha256_file(staged) != record.sha256:
                            raise ArchiveError(f"Diagnostic changed while archiving: {record.source_path}")
                        transaction.replace(staged, record.archive_path)
                    warnings = transaction.commit(
                        cleanup=[record.source_path for record in records],
                        expected_cleanup={record.source_path: record.sha256 for record in records},
                    )
            for folder in (review_xlsx.parent, legacy_review_folder):
                try:
                    folder.rmdir()
                except OSError:
                    pass
            return warnings
    except (ArchiveError, OSError, ValueError, RuntimeError) as exc:
        return [f"Review cleanup pending: {exc}"]
