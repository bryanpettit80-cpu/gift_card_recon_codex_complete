"""Editable monetary variance narratives carried inside the monthly workbook."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Mapping

from gift_card_recon.close_assessment import ControlOutcome
from gift_card_recon.store_config import get_store_config
from gift_card_recon.utils import sha256_file


SHEET_NAME = "Monthly Variance Explanations"
IDENTITY_SHEET = "_monthly_variance_identity"
MAX_TEXT_LENGTH = 1000
_PREFIXES = ("summary_activity_", "period_pos_", "period_tender_")


def _eligible(control: ControlOutcome) -> bool:
    return (
        control.variance is not None
        and control.variance.is_finite()
        and control.variance != 0
        and (control.code == "darden_summary_match" or control.code.startswith(_PREFIXES))
    )


def _plain_text(value: object) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError("Monthly variance explanations must be plain text.")
    value = value.strip()
    if len(value) > MAX_TEXT_LENGTH:
        raise ValueError(f"Monthly variance explanations cannot exceed {MAX_TEXT_LENGTH} characters.")
    if value.startswith(("=", "+", "-", "@")):
        raise ValueError("Monthly variance explanations cannot be Excel formulas.")
    if any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise ValueError("Monthly variance explanation contains an invalid control character.")
    return value


def write_monthly_variance_explanations_sheet(
    workbook, *, store: str, period: str, controls: Iterable[ControlOutcome]
) -> None:
    """Add protected control amounts and editable reasons to an existing report."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = [control for control in controls if _eligible(control)]
    if not rows:
        return
    if SHEET_NAME in workbook.sheetnames or IDENTITY_SHEET in workbook.sheetnames:
        raise ValueError("Monthly variance explanation sheets already exist.")
    if len({control.code for control in rows}) != len(rows):
        raise ValueError("Monthly variance control codes must be unique.")
    sheet = workbook.create_sheet(SHEET_NAME)
    identity = workbook.create_sheet(IDENTITY_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:L2")
    sheet["A1"] = "MONTHLY VARIANCE EXPLANATIONS"
    sheet["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 20
    sheet["A3"] = "Store"
    sheet["B3"] = str(store)
    sheet["F3"] = "Period"
    sheet["G3"] = period
    sheet.merge_cells("A4:L5")
    sheet["A4"] = (
        "Enter an explanation for each monthly difference in the yellow cells, save, "
        "and rerun Monthly Close. Explained amounts remain visible and can close with review. "
        "Complete required weekly explanations in the weekly reports too."
    )
    sheet["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet["A4"].font = Font(name="Arial", size=10)
    sheet.row_dimensions[4].height = 24
    sheet.row_dimensions[5].height = 20
    sheet.merge_cells("A6:C6")
    sheet.merge_cells("F6:L6")
    for cell, text in (("A6", "Control"), ("D6", "Variance"), ("E6", "Status"), ("F6", "Explanation")):
        sheet[cell] = text
    border = Border(bottom=Side(style="thin", color="B7C9DB"))
    for row in sheet.iter_rows(min_row=6, max_row=6, max_col=12):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[6].height = 27
    identity.append(["schema", 1])
    identity.append(["store", str(store)])
    identity.append(["period", period])
    identity.append(["input_sheet", SHEET_NAME])
    identity.append(["code", "label", "variance", "input_row"])
    validation = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(MAX_TEXT_LENGTH), allow_blank=True)
    validation.showErrorMessage = True
    validation.error = f"Use plain text, up to {MAX_TEXT_LENGTH} characters."
    validation.errorTitle = "Explanation too long"
    sheet.add_data_validation(validation)
    for row_number, control in enumerate(rows, start=7):
        explanation = _plain_text(getattr(control, "explanation", ""))
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=3)
        sheet.merge_cells(start_row=row_number, start_column=6, end_row=row_number, end_column=12)
        sheet.cell(row_number, 1, control.label)
        sheet.cell(row_number, 4, float(control.variance))
        sheet.cell(row_number, 4).number_format = '$#,##0.00;($#,##0.00);$0.00'
        sheet.cell(row_number, 5, control.disposition.value)
        reason = sheet.cell(row_number, 6, explanation or None)
        reason.fill = PatternFill("solid", fgColor="FFF2CC")
        reason.protection = Protection(locked=False)
        validation.add(reason)
        for cells in sheet.iter_rows(min_row=row_number, max_row=row_number, max_col=12):
            for cell in cells:
                cell.font = Font(name="Arial", size=10, color="333333")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
        # Sixteen lines accommodate a 1,000-character explanation at normal size.
        sheet.row_dimensions[row_number].height = min(400, max(240, 15 * (2 + max(len(explanation) // 80, explanation.count("\n") + 1))))
        identity.append([control.code, control.label, str(control.variance), row_number])
    for column in "ABCFGHIJKL":
        sheet.column_dimensions[column].width = 11
    sheet.column_dimensions["D"].width = 15
    sheet.column_dimensions["E"].width = 15
    sheet.freeze_panes = "F7"
    sheet.protection.sheet = True
    sheet.protection.selectLockedCells = True
    sheet.protection.selectUnlockedCells = False
    sheet.print_title_rows = "1:6"
    sheet.print_area = f"A1:L{6 + len(rows)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    identity.sheet_state = "veryHidden"


def locate_monthly_variance_explanation_source(
    *, input_dir: Path, store: str, period: str, output_root: Path | None = None
) -> Path | None:
    """Locate the current editable review, falling back to a published close."""
    config = get_store_config(store)
    input_dir = Path(input_dir)
    if output_root is None:
        monthly_root = next((parent for parent in input_dir.parents if parent.name == "02 Monthly Close Inputs"), None)
        if monthly_root is None:
            return None
        output_root = monthly_root.parent / "03 Finished Reports"
    output_root = Path(output_root)
    review = output_root / "Monthly Close - Review Required" / f"{config.output_slug}_{period}_Review_Required.xlsx"
    if review.is_file():
        return review
    from gift_card_recon.fiscal_calendar import fiscal_period_for_label
    fiscal_period = fiscal_period_for_label(period)
    canonical = output_root / "Monthly Close" / fiscal_period.folder_name / f"{config.output_slug}_{period}_Monthly_Close.xlsx"
    return canonical if canonical.is_file() else None


def read_monthly_variance_explanations(
    path: Path, *, store: str, period: str, controls: Iterable[ControlOutcome]
) -> Mapping[str, str]:
    """Accept only plain text bound to the same store, period and current amount."""
    from openpyxl import load_workbook

    path = Path(path)
    before = sha256_file(path)
    workbook = load_workbook(path, data_only=False, keep_links=False)
    try:
        if SHEET_NAME not in workbook.sheetnames and IDENTITY_SHEET not in workbook.sheetnames:
            return {}
        if SHEET_NAME not in workbook.sheetnames or IDENTITY_SHEET not in workbook.sheetnames:
            raise ValueError("Monthly explanation sheet or identity metadata is missing.")
        sheet, identity = workbook[SHEET_NAME], workbook[IDENTITY_SHEET]
        if identity.sheet_state != "veryHidden":
            raise ValueError("Monthly explanation identity metadata must remain very hidden.")
        for worksheet in (sheet, identity):
            if any(cell.data_type == "f" for row in worksheet for cell in row):
                raise ValueError("Monthly explanation inputs and identity cannot contain formulas.")
        if (
            identity["B1"].value != 1
            or str(identity["B2"].value) != str(store)
            or identity["B3"].value != period
            or identity["B4"].value != SHEET_NAME
            or str(sheet["B3"].value) != str(store)
            or sheet["G3"].value != period
        ):
            raise ValueError("Monthly explanation store, period, or schema identity does not match.")
        current = {control.code: control for control in controls if _eligible(control)}
        result: dict[str, str] = {}
        seen_codes: set[str] = set()
        seen_rows: set[int] = set()
        for code, label, amount, row_number in identity.iter_rows(min_row=6, max_col=4, values_only=True):
            if not isinstance(code, str) or code in seen_codes:
                raise ValueError("Monthly explanation identity contains invalid or duplicate control codes.")
            if not isinstance(row_number, int) or isinstance(row_number, bool) or row_number < 7 or row_number in seen_rows or row_number > sheet.max_row:
                raise ValueError("Monthly explanation identity contains an invalid input row.")
            seen_codes.add(code)
            seen_rows.add(row_number)
            try:
                bound_amount = Decimal(str(amount))
                visible_amount = Decimal(str(sheet.cell(row_number, 4).value))
            except InvalidOperation as exc:
                raise ValueError("Monthly explanation control amount is invalid.") from exc
            if not bound_amount.is_finite() or bound_amount != visible_amount or sheet.cell(row_number, 1).value != label:
                raise ValueError("Monthly explanation displayed controls differ from protected identity.")
            text = _plain_text(sheet.cell(row_number, 6).value)
            if not text:
                continue
            control = current.get(code)
            if control is None or control.variance != bound_amount or control.label != label:
                raise ValueError(f"Monthly explanation for {label!r} does not match current controls; review the changed amount.")
            result[code] = text
        if before != sha256_file(path):
            raise ValueError("Monthly variance explanation workbook changed while being read.")
        return result
    finally:
        workbook.close()


def load_monthly_variance_explanations(
    *, input_dir: Path, store: str, period: str, controls: Iterable[ControlOutcome], output_root: Path | None = None
) -> Mapping[str, str]:
    path = locate_monthly_variance_explanation_source(input_dir=input_dir, store=store, period=period, output_root=output_root)
    return {} if path is None else read_monthly_variance_explanations(path, store=store, period=period, controls=controls)
