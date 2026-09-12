from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
import pytest

from gift_card_recon.monthly_close_service import CloseBlockedError
from gift_card_recon.monthly_explanations import IDENTITY_SHEET, SHEET_NAME
from gift_card_recon.utils import sha256_file
from test_monthly_close_service import _build_period, _fake_pdf_exporter, _run


EXPLANATION = "Corporate adjustment reviewed with accounting; retain this operator explanation."


def _review_with_text(root: Path):
    setup = _build_period(root, store="9355")
    workbook = load_workbook(setup["summary_path"])
    workbook["Summary"]["D3"] = 250
    workbook["Summary"]["H3"] = -50
    workbook.save(setup["summary_path"])
    workbook.close()
    with pytest.raises(CloseBlockedError) as blocked:
        _run(setup)
    review = blocked.value.review_workbook
    workbook = load_workbook(review)
    for row in range(7, workbook[SHEET_NAME].max_row + 1):
        workbook[SHEET_NAME].cell(row, 6, EXPLANATION)
    workbook.save(review)
    workbook.close()
    return setup, review


def _saved_path(review: Path, digest: str) -> Path:
    return review.parent / "Saved Explanation Inputs" / f"inputs_{digest[:12]}.xlsx"


def _assert_retained(review: Path, original: bytes) -> Path:
    import hashlib

    digest = hashlib.sha256(original).hexdigest()
    retained = _saved_path(review, digest)
    assert retained.read_bytes() == original
    assert sha256_file(retained) == digest
    workbook = load_workbook(review)
    report_text = "\n".join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    assert str(retained) in report_text
    return retained


def test_stale_amount_retains_operator_reasons_before_replacing_review(tmp_path: Path) -> None:
    setup, review = _review_with_text(tmp_path)
    original = review.read_bytes()
    workbook = load_workbook(setup["summary_path"])
    workbook["Summary"]["D3"] = 251
    workbook.save(setup["summary_path"])
    workbook.close()

    with pytest.raises(CloseBlockedError, match="does not match current controls") as blocked:
        _run(setup)

    assert blocked.value.review_workbook == review
    retained = _assert_retained(review, original)
    workbook = load_workbook(retained)
    assert workbook[SHEET_NAME]["F7"].value == EXPLANATION
    workbook.close()
    assert not (setup["archive_root"] / "Monthly Close").exists()


def test_unrelated_missing_evidence_preserves_entered_monthly_reasons(tmp_path: Path) -> None:
    setup, review = _review_with_text(tmp_path)
    original = review.read_bytes()
    (setup["micros_dir"] / "TENDER_DETAIL.TXT").unlink()

    with pytest.raises(CloseBlockedError, match="TENDER_DETAIL.TXT"):
        _run(setup)

    _assert_retained(review, original)
    assert not (setup["archive_root"] / "Monthly Close").exists()


@pytest.mark.parametrize("sheet,cell,value", [
    (IDENTITY_SHEET, "B1", 99),
    (SHEET_NAME, "F7", "=1+1"),
])
def test_invalid_identity_or_formula_input_is_retained_without_authorizing_close(
    tmp_path: Path, sheet: str, cell: str, value: object,
) -> None:
    setup, review = _review_with_text(tmp_path)
    workbook = load_workbook(review)
    workbook[sheet][cell] = value
    workbook.save(review)
    workbook.close()
    original = review.read_bytes()

    with pytest.raises(CloseBlockedError):
        _run(setup)

    retained = _assert_retained(review, original)
    workbook = load_workbook(retained, data_only=False)
    assert workbook[sheet][cell].value == value
    workbook.close()
    assert not (setup["archive_root"] / "Monthly Close").exists()


def test_existing_mismatched_retention_copy_stops_without_replacing_original_review(tmp_path: Path) -> None:
    setup, review = _review_with_text(tmp_path)
    original = review.read_bytes()
    original_pdf = review.with_suffix(".pdf").read_bytes()
    retained = _saved_path(review, sha256_file(review))
    retained.parent.mkdir()
    retained.write_bytes(b"different retained content")
    (setup["micros_dir"] / "TENDER_DETAIL.TXT").unlink()

    with pytest.raises(CloseBlockedError) as blocked:
        _run(setup)

    assert "size/SHA-256 verification" in blocked.value.review_publication_error
    assert "TENDER_DETAIL.TXT" in str(blocked.value)
    assert review.read_bytes() == original
    assert review.with_suffix(".pdf").read_bytes() == original_pdf
    assert retained.read_bytes() == b"different retained content"


def test_diagnostic_pair_rollback_keeps_original_review_and_verified_retention(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    setup, review = _review_with_text(tmp_path)
    original = review.read_bytes()
    review_pdf = review.with_suffix(".pdf")
    original_pdf = review_pdf.read_bytes()
    retained = _saved_path(review, sha256_file(review))
    (setup["micros_dir"] / "TENDER_DETAIL.TXT").unlink()
    import gift_card_recon.monthly_close_service as service

    real_replace = service.os.replace

    def fail_new_pdf(source, destination):
        if Path(destination) == review_pdf and not str(source).endswith(".backup"):
            raise OSError("simulated diagnostic PDF publication failure")
        return real_replace(source, destination)

    monkeypatch.setattr(service.os, "replace", fail_new_pdf)
    with pytest.raises(CloseBlockedError) as blocked:
        _run(setup)

    assert "simulated diagnostic PDF publication failure" in blocked.value.review_publication_error
    assert review.read_bytes() == original
    assert review_pdf.read_bytes() == original_pdf
    assert retained.read_bytes() == original


def test_failed_atomic_retention_copy_keeps_existing_diagnostic_pair(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    setup, review = _review_with_text(tmp_path)
    original = review.read_bytes()
    original_pdf = review.with_suffix(".pdf").read_bytes()
    (setup["micros_dir"] / "TENDER_DETAIL.TXT").unlink()
    import gift_card_recon.monthly_close_service as service

    def fail_link(*_args, **_kwargs):
        raise OSError("simulated atomic retention copy failure")

    monkeypatch.setattr(service.os, "link", fail_link)
    with pytest.raises(CloseBlockedError) as blocked:
        _run(setup)

    assert "simulated atomic retention copy failure" in blocked.value.review_publication_error
    assert review.read_bytes() == original
    assert review.with_suffix(".pdf").read_bytes() == original_pdf
    assert list((review.parent / "Saved Explanation Inputs").iterdir()) == []


def test_edit_during_diagnostic_render_is_not_overwritten(tmp_path: Path) -> None:
    setup, review = _review_with_text(tmp_path)
    original = review.read_bytes()
    retained = _saved_path(review, sha256_file(review))
    original_pdf = review.with_suffix(".pdf").read_bytes()
    (setup["micros_dir"] / "TENDER_DETAIL.TXT").unlink()
    latest_text = "A new note was added while the diagnostic was rendering."

    def concurrent_edit(**kwargs):
        workbook = load_workbook(review)
        workbook[SHEET_NAME]["F7"] = latest_text
        workbook.save(review)
        workbook.close()
        return _fake_pdf_exporter(**kwargs)

    with pytest.raises(CloseBlockedError) as blocked:
        _run(setup, pdf_exporter=concurrent_edit)

    assert "changed before diagnostic publication" in blocked.value.review_publication_error
    workbook = load_workbook(review)
    assert workbook[SHEET_NAME]["F7"].value == latest_text
    workbook.close()
    assert review.with_suffix(".pdf").read_bytes() == original_pdf
    assert retained.read_bytes() == original
