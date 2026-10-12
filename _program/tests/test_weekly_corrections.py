from datetime import date
from decimal import Decimal
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from gift_card_recon.auto_run import run_weekly_reconciliations
from gift_card_recon.business_identity import activity_identity
from gift_card_recon.business_identity import CorrectionRequired
from gift_card_recon.parsers import parse_activity_file
from gift_card_recon.weekly_service import restore_weekly_revision
from gift_card_recon.variance_explanation import EXPLANATION_SHEET, EXPLANATION_INPUT_CELL
from test_auto_run import make_layout, create_activity, create_micros_week
import pytest


def setup_week(tmp_path, *, pos_payment="20.00"):
    paths = make_layout(tmp_path)
    activity = paths["input"] / "9355 Virginia Beach" / "activity" / "weekly.xlsx"
    create_activity(activity, store="9355", begin=date(2026, 6, 29), end=date(2026, 7, 5),
                    issue=Decimal("10.00"), payment=Decimal("20.00"))
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 29),
                       issue=Decimal("10.00"), payment=Decimal(pos_payment))
    return paths, activity


def execute(paths, token=None):
    return run_weekly_reconciliations(operations_root=paths["operations"], input_root=paths["input"],
        output_dir=paths["output"], monthly_close_root=paths["monthly"], archive_root=paths["archive"],
        review_root=paths["review"], approved_corrections={"9355": token} if token else None)[0]


def package(paths):
    return paths["archive"] / "9355 Virginia Beach" / "2026" / "2026-W27"


def redeliver(paths, destination):
    manifest = json.loads((package(paths) / "weekly_manifest.json").read_text())
    shutil.copy2(package(paths) / manifest["artifacts"]["archived_activity"]["relative_path"], destination)


def test_resaved_renamed_input_and_plain_report_are_semantic_duplicate(tmp_path):
    paths, activity = setup_week(tmp_path)
    first = execute(paths)
    assert first.status == "created", first.message
    renamed_report = first.output_path.with_name("Saved review.xlsx")
    book = load_workbook(first.output_path)
    book["Reconciliation"]["A1"].font = Font(bold=True, size=18)
    book.properties.lastModifiedBy = "Manager"
    book.save(renamed_report)
    book.close()
    first.output_path.unlink()
    renamed_input = activity.with_name("Another name.xlsx")
    redeliver(paths, renamed_input)
    book = load_workbook(renamed_input)
    book.active.column_dimensions["A"].width = 38
    book.save(renamed_input)
    book.close()
    repeat = execute(paths)
    assert repeat.status == "duplicate", repeat.message
    assert repeat.result_id == first.result_id
    assert repeat.output_path == renamed_report
    assert not renamed_input.exists()


def test_record_multiplicity_and_details_require_correction_even_when_totals_match(tmp_path):
    paths, activity = setup_week(tmp_path)
    first = execute(paths)
    redeliver(paths, activity)
    original = parse_activity_file(activity)
    book = load_workbook(activity)
    row = [cell.value for cell in book.active[5]]
    book.active["G5"] = 5
    row[6] = 5
    book.active.append(row)
    book.save(activity)
    book.close()
    assert activity_identity(parse_activity_file(activity)) != activity_identity(original)
    proposal = execute(paths)
    assert proposal.status == "correction-required", proposal.message
    assert "activity_records" in proposal.message
    assert "3" in proposal.message
    assert first.output_path.exists()


def test_repeated_corrections_preserve_previous_versions_and_reject_stale_approval(tmp_path):
    paths, activity = setup_week(tmp_path)
    first = execute(paths)
    before = first.output_path.read_bytes()
    create_activity(activity, store="9355", begin=date(2026, 6, 29), end=date(2026, 7, 5),
                    issue=Decimal("11.00"), payment=Decimal("20.00"))
    proposal = execute(paths)
    assert proposal.status == "correction-required", proposal.message
    assert first.output_path.read_bytes() == before
    corrected = execute(paths, proposal.correction_token)
    assert corrected.status == "created", corrected.message
    manifest = json.loads((package(paths) / "weekly_manifest.json").read_text())
    assert manifest["revision"] == 2
    prior = next(package(paths).with_name("2026-W27.revisions").iterdir())
    assert (prior / "operator_workbook.xlsx").read_bytes() == before
    create_activity(activity, store="9355", begin=date(2026, 6, 29), end=date(2026, 7, 5),
                    issue=Decimal("12.00"), payment=Decimal("20.00"))
    stale = execute(paths, proposal.correction_token)
    assert stale.status == "correction-required", stale.message
    assert stale.correction_token != proposal.correction_token
    third = execute(paths, stale.correction_token)
    assert third.status == "created", third.message
    assert len(list(package(paths).with_name("2026-W27.revisions").iterdir())) == 2
    assert json.loads((package(paths) / "weekly_manifest.json").read_text())["revision"] == 3


def test_pos_only_change_requires_bound_replacement_decision(tmp_path):
    paths, activity = setup_week(tmp_path)
    execute(paths)
    redeliver(paths, activity)
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 29),
                       issue=Decimal("11.00"), payment=Decimal("20.00"))
    proposal = execute(paths)
    assert proposal.status == "correction-required", proposal.message
    assert "pos_gift_card_issue" in proposal.message


def test_corrected_amounts_do_not_reuse_stale_explanation(tmp_path):
    paths, activity = setup_week(tmp_path, pos_payment="40.00")
    first = execute(paths)
    book = load_workbook(first.output_path)
    book[EXPLANATION_SHEET][EXPLANATION_INPUT_CELL] = "Applies to old amount only."
    book.save(first.output_path)
    book.close()

    create_activity(activity, store="9355", begin=date(2026, 6, 29), end=date(2026, 7, 5),
                    issue=Decimal("10.00"), payment=Decimal("30.00"))
    proposal = execute(paths)
    corrected = execute(paths, proposal.correction_token)
    assert corrected.status == "created", corrected.message
    book = load_workbook(corrected.output_path)
    assert book[EXPLANATION_SHEET][EXPLANATION_INPUT_CELL].value in (None, "")
    book.close()
    previous = next(package(paths).with_name("2026-W27.revisions").iterdir())
    book = load_workbook(previous / "operator_workbook.xlsx")
    assert book[EXPLANATION_SHEET][EXPLANATION_INPUT_CELL].value == "Applies to old amount only."
    book.close()


def test_select_previous_weekly_revision_retains_current_and_restores_inputs(tmp_path):
    paths, activity = setup_week(tmp_path)
    first = execute(paths)
    create_activity(activity, store="9355", begin=date(2026, 6, 29), end=date(2026, 7, 5),
                    issue=Decimal("11.00"), payment=Decimal("20.00"))
    proposal = execute(paths)
    current = execute(paths, proposal.correction_token)
    assert current.status == "created", current.message
    selected = next(package(paths).with_name("2026-W27.revisions").iterdir())
    args = dict(store="9355", period="2026-W27", revision_id=selected.name,
                output_root=paths["output"], archive_root=paths["archive"], monthly_close_root=paths["monthly"])
    with pytest.raises(CorrectionRequired) as needed:
        restore_weekly_revision(**args)
    restored = restore_weekly_revision(**args, correction_token=needed.value.token)
    assert restored.revision == 3
    assert restored.input_set_id == first.input_set_id
    assert restored.result_id != first.result_id
    assert len(list(package(paths).with_name("2026-W27.revisions").iterdir())) == 2
    redeliver(paths, activity)
    # Live controls still match the original after selecting its prior revision.
    duplicate = execute(paths)
    assert duplicate.status == "duplicate", duplicate.message


def test_renamed_embedded_explanation_resolves_by_store_and_week(tmp_path):
    from gift_card_recon.variance_explanation import resolve_live_weekly_explanation_path
    paths, _ = setup_week(tmp_path, pos_payment="40.00")
    result = execute(paths)
    assert result.status == "created", result.message
    renamed = result.output_path.with_name("Manager reviewed.xlsx")
    result.output_path.rename(renamed)
    period_input = paths["monthly"] / "9355 Virginia Beach" / "FY27 M01 - Fiscal June"
    found = resolve_live_weekly_explanation_path(period_input, "9355", date(2026, 7, 5),
        archive_root=paths["archive"], output_root=paths["output"])
    assert found == renamed


def test_renamed_legacy_explanation_resolves_and_duplicate_identity_blocks(tmp_path):
    from gift_card_recon.variance_explanation import (
        WeeklyVarianceExplanation, find_legacy_variance_explanation, write_variance_explanation_workbook)
    directory = tmp_path / "Variance Explanations"
    expected = directory / "Manager note.xlsx"
    explanation = WeeklyVarianceExplanation("9355", date(2026, 6, 29), date(2026, 7, 5),
        Decimal("0"), Decimal("20"), Decimal("-20"), Decimal("0"), "Current explanation")
    write_variance_explanation_workbook(explanation, expected)
    assert find_legacy_variance_explanation(tmp_path, "9355", date(2026, 7, 5)) == expected
    shutil.copy2(expected, directory / "Another name.xlsx")
    with pytest.raises(ValueError, match="Multiple weekly"):
        find_legacy_variance_explanation(tmp_path, "9355", date(2026, 7, 5))


@pytest.mark.parametrize("operation", ["correction", "restore"])
@pytest.mark.parametrize("changed_target", ["report", "monthly_activity"])
def test_late_operator_save_is_retained_and_publication_rolls_back(
    tmp_path, monkeypatch, operation, changed_target,
):
    from gift_card_recon.publication import PublicationTransaction, RecoveryRequired, file_state

    paths, activity = setup_week(tmp_path, pos_payment="40.00")
    current = execute(paths)
    create_activity(activity, store="9355", begin=date(2026, 6, 29), end=date(2026, 7, 5),
                    issue=Decimal("10.00"), payment=Decimal("30.00"))
    proposal = execute(paths)
    if operation == "restore":
        current = execute(paths, proposal.correction_token)
        selected = next(package(paths).with_name("2026-W27.revisions").iterdir())
        args = dict(store="9355", period="2026-W27", revision_id=selected.name,
                    output_root=paths["output"], archive_root=paths["archive"], monthly_close_root=paths["monthly"])
        with pytest.raises(CorrectionRequired) as needed:
            restore_weekly_revision(**args)

    manifest = json.loads((package(paths) / "weekly_manifest.json").read_text())
    monthly_path = Path(manifest["artifacts"]["monthly_staged_activity"]["path"])
    target = current.output_path if changed_target == "report" else monthly_path
    before_package = file_state(package(paths))
    before_output = file_state(current.output_path)
    before_monthly = file_state(monthly_path)
    original_replace = PublicationTransaction.replace
    changed_bytes = []

    def save_just_before_replacement(transaction, staged, destination, **kwargs):
        if Path(destination) == target and not changed_bytes:
            book = load_workbook(target)
            if changed_target == "report":
                book[EXPLANATION_SHEET][EXPLANATION_INPUT_CELL] = "Saved after the previous revision was captured."
            else:
                book.active["G5"] = 999
            book.save(target)
            book.close()
            changed_bytes.append(target.read_bytes())
        return original_replace(transaction, staged, destination, **kwargs)

    monkeypatch.setattr(PublicationTransaction, "replace", save_just_before_replacement)
    if operation == "correction":
        stopped = execute(paths, proposal.correction_token)
        assert stopped.status == "skipped", stopped.message
        assert "changed since validation" in stopped.message
        assert activity.exists()
    else:
        with pytest.raises(RecoveryRequired, match="changed since validation"):
            restore_weekly_revision(**args, correction_token=needed.value.token)
    assert changed_bytes and target.read_bytes() == changed_bytes[0]
    assert file_state(package(paths)) == before_package
    if changed_target == "report":
        assert file_state(monthly_path) == before_monthly
    else:
        assert file_state(current.output_path) == before_output
    expected_revision_count = 1 if operation == "restore" else 0
    assert len(list(package(paths).with_name("2026-W27.revisions").glob("r*-*"))) == expected_revision_count


@pytest.mark.parametrize("cleanup_error", [OSError, RuntimeError])
def test_committed_cleanup_failure_reports_created_with_recovery_pending(tmp_path, monkeypatch, cleanup_error):
    from gift_card_recon import publication

    paths, activity = setup_week(tmp_path)
    with monkeypatch.context() as patch:
        def fail_cleanup(*args):
            raise cleanup_error("simulated journal storage failure after commit")
        patch.setattr(publication, "_finish_cleanup", fail_cleanup)
        result = execute(paths)
    assert result.status == "created", result.message
    assert "committed" in result.message.lower() and "Rollback" not in result.message
    assert "simulated journal storage failure" in result.recovery_action
    assert "transaction.json" in result.recovery_action
    manifest = json.loads((package(paths) / "weekly_manifest.json").read_text())
    assert result.result_id == manifest["result_id"]
    assert result.output_path.is_file()
    assert activity.exists()
    support = publication.support_root_for(paths["archive"])
    with publication.writer_lock(support):
        assert publication.recover_pending(support)
    assert result.output_path.is_file() and not activity.exists()
