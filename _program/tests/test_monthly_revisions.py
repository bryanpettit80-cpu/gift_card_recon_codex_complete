from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font

from gift_card_recon.business_identity import CorrectionRequired, workbook_values_identity
from gift_card_recon.evidence_archive import ArchiveError
from gift_card_recon.monthly_close_service import CloseBlockedError, restore_monthly_revision, retry_monthly_pdf
from gift_card_recon.monthly_close_service import _publish_diagnostic
from gift_card_recon.publication import PublicationTransaction
from gift_card_recon.utils import sha256_file
from test_monthly_close_service import _build_period, _run, _fake_pdf_exporter


def test_monthly_renames_formatting_corrections_conflicts_and_restore(tmp_path: Path) -> None:
    setup = _build_period(tmp_path, store="9355")
    first = _run(setup)
    original_values = workbook_values_identity(first.workbook_path)
    original_manifest = sha256_file(first.manifest_path)

    summary = setup["summary_path"]
    book = load_workbook(summary)
    book.active["J1"] = "Report End"
    book.active["K1"] = setup["period"].end_date
    book.save(summary)
    book.close()
    summary.rename(summary.with_name("corporate summary renamed.xlsx"))
    for index, activity in enumerate(sorted((setup["input_dir"] / "activity").glob("*.xlsx"))):
        activity.rename(activity.with_name(f"week {index + 1} renamed.xlsx"))
    book = load_workbook(first.workbook_path)
    book.active["A1"].font = Font(name="Arial", bold=True, color="123456")
    book.active.column_dimensions["A"].width = 27
    book.properties.creator = "Another Windows computer"
    book.save(first.workbook_path)
    book.close()
    renamed_report = first.workbook_path.with_name("Reviewed monthly workbook.xlsx")
    first.workbook_path.rename(renamed_report)
    repeated = _run(setup)
    assert repeated.disposition == "no_change"
    assert repeated.result_id == first.result_id
    assert repeated.input_set_id == first.input_set_id
    assert repeated.run_id != first.run_id
    assert repeated.workbook_path == renamed_report
    assert sha256_file(first.manifest_path) == original_manifest

    activity = sorted((setup["input_dir"] / "activity").glob("*.xlsx"))[0]
    book = load_workbook(activity)
    book.active["E3"] = 500
    book.save(activity)
    book.close()
    with pytest.raises(CorrectionRequired) as proposed:
        _run(setup)
    assert "records" in proposed.value.changes
    assert sha256_file(first.manifest_path) == original_manifest
    book = load_workbook(activity)
    book.active["E3"] = 501
    book.save(activity)
    book.close()
    with pytest.raises(CorrectionRequired) as superseded:
        _run(setup, approve_correction=proposed.value.token)
    assert superseded.value.token != proposed.value.token
    corrected = _run(setup, approve_correction=superseded.value.token)
    assert corrected.result_id != first.result_id
    revisions = corrected.manifest_path.parent / "revisions"
    assert (revisions / first.result_id[:16] / "workbook.xlsx").is_file()
    assert (revisions / corrected.result_id[:16] / "workbook.xlsx").is_file()
    with pytest.raises(CorrectionRequired) as restoration:
        restore_monthly_revision(archive_root=setup["archive_root"], store="9355",
                                 period=setup["period"].period_key, revision_id=first.result_id)
    restored = restore_monthly_revision(
        archive_root=setup["archive_root"], store="9355", period=setup["period"].period_key,
        revision_id=first.result_id, approve_correction=restoration.value.token,
    )
    assert workbook_values_identity(restored) == original_values
    manifest = json.loads(corrected.manifest_path.read_text(encoding="utf-8"))
    assert manifest["restored_from_result_id"] == first.result_id
    assert manifest["previous_result_id"] == corrected.result_id


def test_monthly_changed_report_values_are_not_a_harmless_save(tmp_path: Path) -> None:
    setup = _build_period(tmp_path, store="9355")
    result = _run(setup)
    book = load_workbook(result.workbook_path)
    book["Monthly Close Report"]["A1"] = "Different restaurant"
    book.save(result.workbook_path)
    book.close()
    with pytest.raises((ArchiveError, CorrectionRequired)) as blocked:
        _run(setup)
    if isinstance(blocked.value, CorrectionRequired):
        with pytest.raises(ArchiveError, match="values or formulas changed"):
            _run(setup, approve_correction=blocked.value.token)


def test_monthly_correction_retains_operator_save_during_publication(tmp_path: Path, monkeypatch) -> None:
    setup = _build_period(tmp_path, store="9355")
    first = _run(setup)
    prior_manifest = first.manifest_path.read_bytes()
    activity = sorted((setup["input_dir"] / "activity").glob("*.xlsx"))[0]
    book = load_workbook(activity)
    book.active["E3"] = 500
    book.save(activity)
    book.close()
    with pytest.raises(CorrectionRequired) as proposed:
        _run(setup)
    real_replace = PublicationTransaction.replace
    saved = []

    def save_before_replace(transaction, source, destination, **kwargs):
        if Path(destination) == first.workbook_path:
            book = load_workbook(first.workbook_path)
            book.active["A1"].font = Font(name="Arial", bold=True, color="ABC123")
            book.save(first.workbook_path)
            book.close()
            saved.append(first.workbook_path.read_bytes())
        return real_replace(transaction, source, destination, **kwargs)

    monkeypatch.setattr(PublicationTransaction, "replace", save_before_replace)
    with pytest.raises(CloseBlockedError, match="changed since validation"):
        _run(setup, approve_correction=proposed.value.token)
    assert len(saved) == 1
    assert first.workbook_path.read_bytes() == saved[0]
    assert first.manifest_path.read_bytes() == prior_manifest
    revisions = first.manifest_path.parent / "revisions"
    assert [path.name for path in revisions.iterdir()] == [first.result_id[:16]]


def test_monthly_pdf_retry_rejects_workbook_save_during_export(tmp_path: Path) -> None:
    setup = _build_period(tmp_path, store="9355")
    first = _run(setup)
    prior_manifest = first.manifest_path.read_bytes()
    prior_pdf = first.pdf_path.read_bytes()

    def export_after_save(**kwargs):
        book = load_workbook(kwargs["workbook_path"])
        book.active["A1"].font = Font(name="Arial", bold=True, color="ABC123")
        book.save(kwargs["workbook_path"])
        book.close()
        return _fake_pdf_exporter(**kwargs)

    with pytest.raises(ArchiveError, match="changed since validation"):
        retry_monthly_pdf(archive_root=setup["archive_root"], store="9355",
                          period=setup["period"].period_key, pdf_exporter=export_after_save)
    assert first.manifest_path.read_bytes() == prior_manifest
    assert first.pdf_path.read_bytes() == prior_pdf


def test_monthly_duplicate_returns_verified_manifest_evidence_after_input_resave(tmp_path: Path) -> None:
    setup = _build_period(tmp_path, store="9355")
    first = _run(setup)
    prior_sources = {record.archive_path.resolve(): record.sha256 for record in first.archive_records}
    input_hash = sha256_file(setup["summary_path"])
    book = load_workbook(setup["summary_path"])
    book.active["D3"].font = Font(name="Arial", bold=True, color="128800")
    book.save(setup["summary_path"])
    book.close()
    assert sha256_file(setup["summary_path"]) != input_hash
    repeated = _run(setup)
    assert repeated.disposition == "no_change"
    assert {record.archive_path.resolve(): record.sha256 for record in repeated.archive_records} == prior_sources
    assert all(record.archive_path.is_file() and sha256_file(record.archive_path) == record.sha256
               for record in repeated.archive_records)
    assert all(not record.remove_after_publish for record in repeated.archive_records)
    preserved_summary = next(record.archive_path for record in repeated.archive_records if record.role == "Gift Card Summary")
    preserved_summary.write_bytes(b"changed archived evidence")
    with pytest.raises(ArchiveError, match="Preserved monthly evidence no longer verifies"):
        _run(setup)


@pytest.mark.parametrize("failure", ["second_artifact", "changed_staged_values"])
def test_diagnostic_uses_durable_rollback_and_attestation(tmp_path: Path, monkeypatch, failure: str) -> None:
    output = tmp_path / "Output" / "Monthly Close - Review Required"
    output.mkdir(parents=True)
    workbook, pdf = output / "review.xlsx", output / "review.pdf"
    workbook.write_bytes(b"previous workbook")
    pdf.write_bytes(b"previous PDF")
    staged_workbook, staged_pdf = tmp_path / "new.xlsx", tmp_path / "new.pdf"
    staged_workbook.write_bytes(b"new workbook")
    staged_pdf.write_bytes(b"new PDF")
    original = PublicationTransaction.replace

    def fail(transaction, source, destination):
        if Path(destination) == pdf and failure == "second_artifact":
            raise OSError("simulated unavailable storage")
        if Path(destination) == workbook and failure == "changed_staged_values":
            Path(source).write_bytes(b"different values after attestation")
        return original(transaction, source, destination)

    monkeypatch.setattr(PublicationTransaction, "replace", fail)
    with pytest.raises(OSError):
        _publish_diagnostic(temp_xlsx=staged_workbook, temp_pdf=staged_pdf,
                            canonical_xlsx=workbook, canonical_pdf=pdf)
    assert workbook.read_bytes() == b"previous workbook"
    assert pdf.read_bytes() == b"previous PDF"
    journals = list((tmp_path / "_automation_runs" / "support" / "transactions").glob("*/transaction.json"))
    assert journals
    assert all(json.loads(path.read_text(encoding="utf-8"))["state"] == "rolled_back" for path in journals)


def test_monthly_pdf_retry_resolves_handoff_workspace_and_renamed_workbook(tmp_path: Path) -> None:
    setup = _build_period(tmp_path, store="9355")
    result = _run(setup)
    original_pdf_hash = sha256_file(result.pdf_path)
    other = tmp_path / "Another username and path"
    archive = other / setup["archive_root"].name
    output = other / setup["output_root"].name
    shutil.copytree(setup["archive_root"], archive)
    shutil.copytree(setup["output_root"], output)
    moved_workbook = other / result.workbook_path.relative_to(tmp_path)
    renamed = moved_workbook.with_name("Renamed after synchronization.xlsx")
    moved_workbook.rename(renamed)
    pdf = retry_monthly_pdf(archive_root=archive, store="9355", period=setup["period"].period_key,
                            pdf_exporter=_fake_pdf_exporter)
    assert pdf == renamed.with_suffix(".pdf")
    assert pdf.is_file()
    assert sha256_file(result.pdf_path) == original_pdf_hash
