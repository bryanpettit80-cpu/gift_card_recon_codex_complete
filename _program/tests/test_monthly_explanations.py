from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from gift_card_recon.close_assessment import ControlDisposition, ControlOutcome
from gift_card_recon.monthly_explanations import (
    IDENTITY_SHEET,
    MAX_TEXT_LENGTH,
    SHEET_NAME,
    load_monthly_variance_explanations,
    locate_monthly_variance_explanation_source,
    read_monthly_variance_explanations,
    write_monthly_variance_explanations_sheet,
)


STORE = "9355"
PERIOD = "FY27-M03"
CONTROL = ControlOutcome(
    code="period_pos_payment",
    label="Period POS payment",
    disposition=ControlDisposition.BLOCK,
    message="POS payment differs from activity.",
    variance=Decimal("35.00"),
)


def _write_report(path: Path, *, text: str = "", controls: tuple[ControlOutcome, ...] | None = None) -> Path:
    workbook = Workbook()
    workbook.active.title = "Monthly Close"
    workbook.active["A1"] = "=SUM(1,2)"
    write_monthly_variance_explanations_sheet(
        workbook, store=STORE, period=PERIOD,
        controls=controls or (replace(CONTROL, explanation=text),),
    )
    workbook.save(path)
    workbook.close()
    return path


def _read(path: Path, *, controls: tuple[ControlOutcome, ...] = (CONTROL,)):
    return read_monthly_variance_explanations(path, store=STORE, period=PERIOD, controls=controls)


def _edit(path: Path, sheet: str, coordinate: str, value: object) -> None:
    workbook = load_workbook(path)
    workbook[sheet][coordinate] = value
    workbook.save(path)
    workbook.close()


def test_monthly_explanations_round_trip_plain_text_with_normal_report_formulas(tmp_path: Path) -> None:
    text = "  Payment recorded in Micros on August 27.\nTerminal redemption requires follow-up.  "
    path = _write_report(tmp_path / "report.xlsx", text=text)
    assert _read(path) == {CONTROL.code: text.strip()}
    workbook = load_workbook(path)
    assert workbook["Monthly Close"]["A1"].data_type == "f"
    assert workbook[SHEET_NAME]["F7"].protection.locked is False
    assert workbook[SHEET_NAME]["D7"].protection.locked is True
    assert workbook[SHEET_NAME].protection.sheet is True
    assert workbook[IDENTITY_SHEET].sheet_state == "veryHidden"
    workbook.close()


def test_blank_explanations_do_not_authorize_variances(tmp_path: Path) -> None:
    path = _write_report(tmp_path / "report.xlsx")
    assert _read(path) == {}
    # Old blank placeholders do not block a corrected or newly recalculated amount.
    assert _read(path, controls=(replace(CONTROL, variance=Decimal("40.00")),)) == {}


def test_reports_without_monthly_input_remain_compatible(tmp_path: Path) -> None:
    path = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    workbook.save(path)
    workbook.close()
    assert _read(path) == {}


@pytest.mark.parametrize("sheet", [SHEET_NAME, IDENTITY_SHEET])
def test_missing_one_monthly_explanation_sheet_is_rejected(tmp_path: Path, sheet: str) -> None:
    path = _write_report(tmp_path / "report.xlsx", text="Reviewed.")
    workbook = load_workbook(path)
    del workbook[sheet]
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="sheet or identity metadata is missing"):
        _read(path)


@pytest.mark.parametrize("sheet,cell", [(SHEET_NAME, "F7"), (SHEET_NAME, "D7"), (IDENTITY_SHEET, "C6")])
def test_formula_inputs_and_identity_are_rejected(tmp_path: Path, sheet: str, cell: str) -> None:
    path = _write_report(tmp_path / "report.xlsx")
    _edit(path, sheet, cell, "=1+1")
    with pytest.raises(ValueError, match="cannot contain formulas"):
        _read(path)


@pytest.mark.parametrize("value", [123, "x" * (MAX_TEXT_LENGTH + 1), " =1+1", "@SUM(1,2)"])
def test_invalid_operator_text_is_rejected(tmp_path: Path, value: object) -> None:
    path = _write_report(tmp_path / "report.xlsx")
    _edit(path, SHEET_NAME, "F7", value)
    with pytest.raises(ValueError, match="plain text|cannot exceed|cannot be Excel formulas|cannot contain formulas"):
        _read(path)


def test_maximum_length_operator_text_is_accepted(tmp_path: Path) -> None:
    text = "x" * MAX_TEXT_LENGTH
    path = _write_report(tmp_path / "report.xlsx", text=text)
    assert _read(path) == {CONTROL.code: text}


@pytest.mark.parametrize(
    "sheet,cell,value",
    [
        (SHEET_NAME, "B3", "9354"),
        (SHEET_NAME, "G3", "FY27-M02"),
        (IDENTITY_SHEET, "B1", 99),
        (IDENTITY_SHEET, "B2", "9354"),
        (IDENTITY_SHEET, "B3", "FY27-M02"),
        (IDENTITY_SHEET, "B4", "Wrong Sheet"),
    ],
)
def test_wrong_store_period_or_schema_cannot_supply_explanations(
    tmp_path: Path, sheet: str, cell: str, value: object,
) -> None:
    path = _write_report(tmp_path / "report.xlsx", text="Reviewed.")
    _edit(path, sheet, cell, value)
    with pytest.raises(ValueError, match="store, period, or schema identity does not match"):
        _read(path)


@pytest.mark.parametrize("state", ["visible", "hidden"])
def test_monthly_identity_must_remain_very_hidden(tmp_path: Path, state: str) -> None:
    path = _write_report(tmp_path / "report.xlsx", text="Reviewed.")
    workbook = load_workbook(path)
    workbook[IDENTITY_SHEET].sheet_state = state
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="very hidden"):
        _read(path)


@pytest.mark.parametrize("control", [
    replace(CONTROL, variance=Decimal("36.00")),
    replace(CONTROL, variance=Decimal("0.00")),
    replace(CONTROL, label="Changed label"),
    replace(CONTROL, code="period_pos_different"),
])
def test_stale_monthly_explanation_is_rejected_when_current_control_changes(
    tmp_path: Path, control: ControlOutcome,
) -> None:
    path = _write_report(tmp_path / "report.xlsx", text="Reviewed.")
    with pytest.raises(ValueError, match="does not match current controls"):
        _read(path, controls=(control,))


@pytest.mark.parametrize("sheet,cell,value", [
    (SHEET_NAME, "D7", 36),
    (SHEET_NAME, "A7", "Changed label"),
    (IDENTITY_SHEET, "C6", "36.00"),
    (IDENTITY_SHEET, "C6", "NaN"),
    (IDENTITY_SHEET, "C6", "not-money"),
    (IDENTITY_SHEET, "D6", True),
    (IDENTITY_SHEET, "D6", 0),
    (IDENTITY_SHEET, "D6", 999),
])
def test_malformed_or_tampered_control_identity_is_rejected(
    tmp_path: Path, sheet: str, cell: str, value: object,
) -> None:
    path = _write_report(tmp_path / "report.xlsx", text="Reviewed.")
    _edit(path, sheet, cell, value)
    with pytest.raises(ValueError, match="displayed controls differ|amount is invalid|invalid input row"):
        _read(path)


@pytest.mark.parametrize("duplicate_code", [True, False])
def test_duplicate_identity_codes_or_input_rows_are_rejected(tmp_path: Path, duplicate_code: bool) -> None:
    path = _write_report(tmp_path / "report.xlsx", text="Reviewed.")
    workbook = load_workbook(path)
    identity = workbook[IDENTITY_SHEET]
    row = [identity.cell(6, column).value for column in range(1, 5)]
    if not duplicate_code:
        row[0] = "period_pos_another"
    identity.append(row)
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="duplicate control codes|invalid input row"):
        _read(path)


def test_writer_omits_zero_and_nonmonetary_integrity_controls(tmp_path: Path) -> None:
    controls = (
        replace(CONTROL, variance=Decimal("0.00")),
        replace(CONTROL, code="activity_identity", variance=None),
        replace(CONTROL, code="darden_identity", variance=Decimal("10.00")),
    )
    workbook = Workbook()
    write_monthly_variance_explanations_sheet(workbook, store=STORE, period=PERIOD, controls=controls)
    assert SHEET_NAME not in workbook.sheetnames
    assert IDENTITY_SHEET not in workbook.sheetnames
    workbook.close()


def test_writer_rejects_duplicate_control_codes() -> None:
    workbook = Workbook()
    with pytest.raises(ValueError, match="control codes must be unique"):
        write_monthly_variance_explanations_sheet(workbook, store=STORE, period=PERIOD, controls=(CONTROL, CONTROL))
    workbook.close()


def test_source_lookup_prefers_current_review_then_published_close(tmp_path: Path) -> None:
    operations = tmp_path / "Gift Card Reconciliation"
    input_dir = operations / "02 Monthly Close Inputs" / "9355 Virginia Beach" / "FY27 M03 - Fiscal August"
    input_dir.mkdir(parents=True)
    output_root = operations / "03 Finished Reports"
    canonical = output_root / "Monthly Close" / "FY27 M03 - Fiscal August" / "Virginia_Beach_9355_FY27-M03_Monthly_Close.xlsx"
    review = output_root / "Monthly Close - Review Required" / "Virginia_Beach_9355_FY27-M03_Review_Required.xlsx"
    assert locate_monthly_variance_explanation_source(input_dir=input_dir, store=STORE, period=PERIOD) is None
    canonical.parent.mkdir(parents=True)
    _write_report(canonical, text="Published explanation.")
    assert locate_monthly_variance_explanation_source(input_dir=input_dir, store=STORE, period=PERIOD) == canonical
    review.parent.mkdir(parents=True)
    _write_report(review, text="Current review explanation.")
    assert locate_monthly_variance_explanation_source(input_dir=input_dir, store=STORE, period=PERIOD) == review
    assert load_monthly_variance_explanations(input_dir=input_dir, store=STORE, period=PERIOD, controls=(CONTROL,)) == {
        CONTROL.code: "Current review explanation.",
    }


def test_lookup_supports_explicit_output_root_for_isolated_inputs(tmp_path: Path) -> None:
    output_root = tmp_path / "reports"
    input_dir = tmp_path / "inputs"
    review = output_root / "Monthly Close - Review Required" / "Virginia_Beach_9355_FY27-M03_Review_Required.xlsx"
    review.parent.mkdir(parents=True)
    _write_report(review)
    assert locate_monthly_variance_explanation_source(input_dir=input_dir, store=STORE, period=PERIOD) is None
    assert locate_monthly_variance_explanation_source(input_dir=input_dir, store=STORE, period=PERIOD, output_root=output_root) == review
