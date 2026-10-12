from __future__ import annotations

import json
import shutil
from datetime import date
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from gift_card_recon.auto_run import run_weekly_reconciliations, weekly_input_inventory
from gift_card_recon.auto_run import main as weekly_main
from gift_card_recon.coordinator import assess_monthly_readiness, assess_weekly_readiness
from gift_card_recon.business_identity import activity_identity, identity, workbook_values_identity
from gift_card_recon.parsers import (ParseError, discover_input_files, parse_activity_file,
    parse_pos_controls, parse_summary, pos_controls_from_args, summary_report_end)
from gift_card_recon.utils import money, sha256_file, to_decimal
from gift_card_recon.business_identity import CorrectionRequired
from gift_card_recon.weekly_service import WeeklyPublication
from test_auto_run import create_activity, create_micros_week, make_layout, run
from test_monthly_close_service import _build_period, _fake_pdf_exporter


def _activity(path: Path, *, store="9355") -> None:
    create_activity(path, store=store, begin=date(2026, 6, 1), end=date(2026, 6, 7),
                    issue=Decimal("50.00"), payment=Decimal("30.00"))


def _change_activity(path: Path, *, amount=50, label="Activation") -> None:
    book = load_workbook(path)
    book.active["C5"] = label
    book.active["G5"] = amount
    book.save(path)
    book.close()


@pytest.mark.parametrize("value", [True, False, "NaN", "sNaN", "Infinity", "-Infinity", float("nan"), float("inf"), Decimal("NaN")])
def test_money_never_interprets_boolean_or_nonfinite_as_currency(value):
    assert to_decimal(value) is None
    with pytest.raises(ValueError, match="finite monetary"):
        money(value)
    with pytest.raises(ParseError, match="POS"):
        pos_controls_from_args("9355", "2026-W23", value, "30.00")


@pytest.mark.parametrize("value", [True, False, "NaN", "Infinity", "1e100000"])
def test_malformed_activity_money_is_a_parse_error(tmp_path, value):
    path = tmp_path / "activity.xlsx"
    _activity(path)
    _change_activity(path, amount=value)
    with pytest.raises(ParseError, match="money value"):
        parse_activity_file(path)


@pytest.mark.parametrize("label", ["Unsupported", "Activation Inquiry", "Redemption Inquiry", "Reload Inquiry", "Page Adjustment", "Void Of Activation Inquiry"])
def test_unknown_transactions_never_disappear(tmp_path, label):
    path = tmp_path / "activity.xlsx"
    _activity(path)
    _change_activity(path, label=label)
    with pytest.raises(ParseError, match="Unsupported transaction type"):
        parse_activity_file(path)


def test_invalid_money_store_does_not_interrupt_valid_store(tmp_path):
    paths = make_layout(tmp_path, stores=("9354 Richmond", "9355 Virginia Beach"))
    bad = paths["input"] / "9354 Richmond" / "activity" / "bad.xlsx"
    good = paths["input"] / "9355 Virginia Beach" / "activity" / "good.xlsx"
    _activity(bad, store="9354")
    _change_activity(bad, amount="Infinity")
    _activity(good)
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 1),
                       issue=Decimal("50"), payment=Decimal("30"))
    reports = run(paths)
    assert [(item.store, item.status) for item in reports] == [("9354", "skipped"), ("9355", "created")]
    assert "money value" in reports[0].message
    assert bad.exists()
    assert reports[1].close_status == "PASS"


def test_summary_and_pos_reject_nonfinite_money(tmp_path):
    setup = _build_period(tmp_path, store="9355")
    book = load_workbook(setup["summary_path"])
    book.active["D3"] = True
    book.save(setup["summary_path"])
    with pytest.raises(ParseError, match="Total Activations"):
        parse_summary(setup["summary_path"], "9355")
    pos = tmp_path / "pos_controls.csv"
    pos.write_text("store,period,pos_gift_card_issue,pos_gift_card_payment\n9355,auto,NaN,30\n", encoding="utf-8")
    with pytest.raises(ParseError, match="malformed value"):
        parse_pos_controls(pos, "9355", "2026-W23")


def test_temporary_files_are_ignored_and_legacy_weekly_folder_selected(tmp_path):
    paths = make_layout(tmp_path, stores=())
    folder = paths["input"] / "9355" / "weekly" / "activity"
    _activity(folder / "renamed.xlsx")
    (folder / "~$renamed.xlsx").write_bytes(b"Excel owner file")
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 1),
                       issue=Decimal("50"), payment=Decimal("30"))
    reports = run(paths)
    assert reports[0].status == "created", reports[0].message


def test_ambiguous_store_does_not_hide_other_store(tmp_path):
    paths = make_layout(tmp_path, stores=("9354 Richmond", "9354 Copy", "9355 Virginia Beach"))
    found, issues = weekly_input_inventory(paths["input"], set())
    assert [store for store, _ in found] == ["9355"]
    assert issues[0][0] == "9354"
    assert "found 2" in issues[0][1]
    assert [(item.store, item.status) for item in run(paths)] == [("9354", "skipped"), ("9355", "no-op")]


def test_missing_activity_folder_is_blocked_not_empty(tmp_path):
    input_root = tmp_path / "inbox"
    (input_root / "9355 Virginia Beach").mkdir(parents=True)
    reports = run_weekly_reconciliations(input_root=input_root, output_dir=tmp_path / "output")
    assert reports[0].status == "skipped"
    assert "folder is missing" in reports[0].message
    assert not (tmp_path / "output").exists()


def test_renamed_monthly_reports_are_discovered_by_content(tmp_path):
    setup = _build_period(tmp_path, store="9355")
    renamed_summary = setup["summary_path"].with_name("renamed workbook.xlsx")
    setup["summary_path"].rename(renamed_summary)
    original_activity = next((setup["input_dir"] / "activity").glob("*.xlsx"))
    original_activity.rename(original_activity.with_name("renamed week.xlsx"))
    (renamed_summary.parent / "~$renamed workbook.xlsx").write_bytes(b"Excel owner")
    summary, activities, pos = discover_input_files(setup["input_dir"])
    assert summary == renamed_summary
    assert len(activities) == 5
    assert any(path.name == "renamed week.xlsx" for path in activities)
    assert pos is None


def test_summary_uses_internal_date_after_rename_and_rejects_conflicts(tmp_path):
    path = tmp_path / "report.xlsx"
    book = Workbook()
    book.active.append(["END DATE: '05-JUL-2026'"])
    book.save(path)
    assert summary_report_end(path) == date(2026, 7, 5)
    book.active.append(["Period ending", "2026-08-02"])
    book.save(path)
    with pytest.raises(ParseError, match="Conflicting"):
        summary_report_end(path)


def test_summary_can_use_byte_bound_intake_metadata(tmp_path):
    path = tmp_path / "renamed.xlsx"
    book = Workbook()
    book.active.append(["SUMMARY"])
    book.save(path)
    (tmp_path / ".gift-card-intake.json").write_text(json.dumps({"files": [
        {"role": "summary", "sha256": sha256_file(path), "report_end": "2026-07-05"}
    ]}), encoding="utf-8")
    assert summary_report_end(path) == date(2026, 7, 5)
    book.active.append(["Changed values"])
    book.save(path)
    assert summary_report_end(path) is None


def test_summary_intake_content_binding_survives_formatting_save(tmp_path):
    path = tmp_path / "renamed.xlsx"
    book = Workbook()
    book.active.append(["SUMMARY"])
    book.save(path)
    (tmp_path / ".gift-card-intake.json").write_text(json.dumps({"files": [
        {"role": "summary", "sha256": sha256_file(path), "content_id": workbook_values_identity(path),
         "report_end": "2026-07-05"}
    ]}), encoding="utf-8")
    book.active.column_dimensions["A"].width = 45
    book.save(path)
    assert summary_report_end(path) == date(2026, 7, 5)
    book.active["A1"] = "Changed values"
    book.save(path)
    assert summary_report_end(path) is None


def test_summary_reads_intake_metadata_from_internal_support_location(tmp_path):
    path = tmp_path / "02 Monthly Close Inputs" / "9355 Virginia Beach" / "FY27 M01 - Fiscal June" / "summary" / "renamed.xlsx"
    path.parent.mkdir(parents=True)
    book = Workbook()
    book.active.append(["SUMMARY"])
    book.save(path)
    metadata = tmp_path / "_automation_runs" / "support" / "intake.json"
    metadata.parent.mkdir(parents=True)
    metadata.write_text(json.dumps({"files": [
        {"role": "summary", "sha256": sha256_file(path), "content_id": workbook_values_identity(path),
         "report_end": "2026-07-05"}
    ]}), encoding="utf-8")
    assert summary_report_end(path) == date(2026, 7, 5)


def _snapshot(root: Path):
    return {str(path.relative_to(root)): (sha256_file(path), path.stat().st_mtime_ns) if path.is_file() else None
            for path in root.rglob("*")}


def test_weekly_readiness_is_repeatable_read_only_and_independent(tmp_path, monkeypatch):
    paths = make_layout(tmp_path, stores=("9354 Richmond", "9355 Virginia Beach"))
    _activity(paths["input"] / "9355 Virginia Beach" / "activity" / "ready.xlsx")
    bad = paths["input"] / "9354 Richmond" / "activity" / "bad.xlsx"
    _activity(bad, store="9354")
    _change_activity(bad, label="Activation Inquiry")
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 1),
                       issue=Decimal("50"), payment=Decimal("30"))
    before = _snapshot(tmp_path)
    def no_mkdir(*args, **kwargs):
        pytest.fail("Readiness attempted to create a directory")
    monkeypatch.setattr(Path, "mkdir", no_mkdir)
    first = assess_weekly_readiness(operations_root=paths["operations"])
    second = assess_weekly_readiness(operations_root=paths["operations"])
    assert first == second
    assert [(item.store, item.ready) for item in first] == [("9354", False), ("9355", True)]
    assert first[1].input_set_id and first[1].result_id
    assert _snapshot(tmp_path) == before


def test_monthly_readiness_uses_same_controls_without_staging(tmp_path, monkeypatch):
    setup = _build_period(tmp_path, store="9355")
    monkeypatch.setattr("gift_card_recon.monthly_close_service.parse_darden_credit_memo", lambda path: setup["darden"])
    before = _snapshot(tmp_path)
    def no_mkdir(*args, **kwargs):
        pytest.fail("Readiness attempted to create a directory")
    monkeypatch.setattr(Path, "mkdir", no_mkdir)
    kwargs = dict(operations_root=tmp_path, period="FY27-M01", stores=["9355"],
                  input_dirs={"9355": setup["input_dir"]}, micros_paths={"9355": setup["micros_dir"]},
                  allow_unconfigured_micros=True)
    first = assess_monthly_readiness(**kwargs)
    assert first == assess_monthly_readiness(**kwargs)
    assert first[0].ready, first[0].reasons
    assert first[0].input_set_id
    assert _snapshot(tmp_path) == before


def test_readiness_reports_pending_recovery_without_replaying(tmp_path):
    paths = make_layout(tmp_path)
    journal = paths["operations"] / "_automation_runs" / "support" / "transactions" / "interrupted" / "transaction.json"
    journal.parent.mkdir(parents=True)
    journal.write_text('{"state":"active"}', encoding="utf-8")
    before = _snapshot(tmp_path)
    result = assess_weekly_readiness(operations_root=paths["operations"], stores=["9355"])[0]
    assert not result.ready
    assert result.recovery_action == "run_recovery"
    assert _snapshot(tmp_path) == before


def test_monthly_readiness_never_extracts_micros_archives(tmp_path, monkeypatch):
    setup = _build_period(tmp_path, store="9355")
    archive = tmp_path / "micros.zip"
    archive.write_bytes(b"not extracted by readiness")
    result = assess_monthly_readiness(operations_root=tmp_path, period="FY27-M01", stores=["9355"],
                                     input_dirs={"9355": setup["input_dir"]}, micros_paths={"9355": archive},
                                     allow_unconfigured_micros=True)[0]
    assert not result.ready
    assert "explicit preparation step" in result.reasons[0]
    assert not (tmp_path / "_automation_runs").exists()


def test_weekly_readiness_reports_duplicate_and_exact_correction_token_without_writes(tmp_path):
    paths = make_layout(tmp_path)
    source = paths["input"] / "9355 Virginia Beach" / "activity" / "week.xlsx"
    _activity(source)
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 1),
                       issue=Decimal("50"), payment=Decimal("30"))
    published = run(paths)[0]
    assert published.status == "created", published.message
    archived = published.archive_path / "activity" / source.name
    shutil.copy2(archived, source)
    before = _snapshot(tmp_path)
    duplicate = assess_weekly_readiness(operations_root=paths["operations"], stores=["9355"])[0]
    assert duplicate.ready and duplicate.disposition == "DUPLICATE"
    assert duplicate.input_set_id == published.input_set_id
    assert _snapshot(tmp_path) == before
    _change_activity(source, amount=51)
    before = _snapshot(tmp_path)
    correction = assess_weekly_readiness(operations_root=paths["operations"], stores=["9355"])[0]
    assert not correction.ready and correction.disposition == "CORRECTION_REQUIRED"
    assert correction.changed_facts and correction.correction_token
    assert _snapshot(tmp_path) == before
    execution = run(paths)[0]
    assert execution.correction_token == correction.correction_token
    assert execution.input_set_id == correction.input_set_id


def test_monthly_readiness_reports_duplicate_and_correction_without_writes(tmp_path, monkeypatch):
    from gift_card_recon.monthly_close_service import run_monthly_close_service
    setup = _build_period(tmp_path, store="9355")
    monkeypatch.setattr("gift_card_recon.monthly_close_service.parse_darden_credit_memo", lambda path: setup["darden"])
    published = run_monthly_close_service(
        store="9355", period="FY27-M01", input_dir=setup["input_dir"], micros_path=setup["micros_dir"],
        micros_work_dir=tmp_path / "extract", archive_root=setup["archive_root"],
        output_root=setup["output_root"], darden_report=setup["darden"], cleanup_sources=False,
        allow_unconfigured_micros=True, pdf_exporter=_fake_pdf_exporter,
    )
    kwargs = dict(operations_root=tmp_path, period="FY27-M01", stores=["9355"],
                  archive_root=setup["archive_root"], output_root=setup["output_root"],
                  input_dirs={"9355": setup["input_dir"]}, micros_paths={"9355": setup["micros_dir"]},
                  allow_unconfigured_micros=True)
    before = _snapshot(tmp_path)
    duplicate = assess_monthly_readiness(**kwargs)[0]
    assert duplicate.ready and duplicate.disposition == "DUPLICATE", duplicate.reasons
    assert duplicate.input_set_id == published.input_set_id
    assert _snapshot(tmp_path) == before
    book = load_workbook(setup["summary_path"])
    book.active["D3"] = 101
    book.save(setup["summary_path"])
    book.close()
    before = _snapshot(tmp_path)
    correction = assess_monthly_readiness(**kwargs)[0]
    assert not correction.ready and correction.disposition == "CORRECTION_REQUIRED", correction.reasons
    assert correction.correction_token and correction.changed_facts
    assert _snapshot(tmp_path) == before


def test_monthly_readiness_requires_correction_for_changed_report_narrative(tmp_path, monkeypatch):
    from openpyxl.styles import Font
    from gift_card_recon import monthly_close_service as service
    from test_monthly_close_service import _run

    setup = _build_period(tmp_path, store="9355")
    book = load_workbook(setup["summary_path"])
    book["Summary"]["D3"], book["Summary"]["H3"] = 250, -50
    book.save(setup["summary_path"])
    book.close()
    with pytest.raises(service.CloseBlockedError) as blocked:
        _run(setup)
    source = blocked.value.review_workbook
    book = load_workbook(source)
    for row in range(7, book["Monthly Variance Explanations"].max_row + 1):
        book["Monthly Variance Explanations"].cell(row, 6, "Reviewed corporate settlement adjustment.")
    book.save(source)
    book.close()
    published = _run(setup)
    retained = next(record for record in published.archive_records if record.role == "monthly_variance_explanation")
    selected_source = service.ArchivedVarianceExplanationSource(
        path=retained.archive_path, sha256=retained.sha256, size_bytes=retained.size_bytes)
    original_prepare = service.prepare_monthly_close_readiness
    monkeypatch.setattr(service, "prepare_monthly_close_readiness", lambda **kwargs: original_prepare(
        **kwargs, archived_variance_explanations={}, archived_monthly_variance_explanation=selected_source))
    archive_input = setup["archive_root"] / "Monthly Close" / "9355" / setup["period"].folder_name
    archived_darden = replace(setup["darden"], source_file=archive_input / "darden" / setup["darden"].source_file.name)
    monkeypatch.setattr(service, "parse_darden_credit_memo", lambda path: archived_darden)
    kwargs = dict(operations_root=tmp_path, period="FY27-M01", stores=["9355"],
                  archive_root=setup["archive_root"], output_root=setup["output_root"],
                  input_dirs={"9355": archive_input}, micros_paths={"9355": archive_input / "micros"},
                  allow_unconfigured_micros=True)

    book = load_workbook(published.workbook_path)
    book["Monthly Close Report"]["A1"].font = Font(bold=True, color="0000FF")
    book.save(published.workbook_path)
    book.close()
    before = _snapshot(tmp_path)
    duplicate = assess_monthly_readiness(**kwargs)[0]
    assert duplicate.ready and duplicate.disposition == "DUPLICATE", duplicate.reasons
    assert duplicate.input_set_id == published.input_set_id
    assert _snapshot(tmp_path) == before

    book = load_workbook(published.workbook_path)
    book["Monthly Variance Explanations"]["F7"] = "A later note is not part of the selected archived proof."
    book.save(published.workbook_path)
    book.close()
    before = _snapshot(tmp_path)
    correction = assess_monthly_readiness(**kwargs)[0]
    assert correction.disposition == "CORRECTION_REQUIRED" and not correction.ready, correction.reasons
    assert correction.input_set_id == published.input_set_id
    assert "report_explanations" in correction.changed_facts
    assert _snapshot(tmp_path) == before
    with pytest.raises(CorrectionRequired) as execution:
        _run(setup, input_dir=archive_input, micros_path=archive_input / "micros",
             darden_report=archived_darden, archived_variance_explanations={},
             archived_monthly_variance_explanation=selected_source)
    assert correction.correction_token == execution.value.token


def test_business_identity_normalizes_equivalent_activity_labels_and_identifier_types(tmp_path):
    path = tmp_path / "activity.xlsx"
    _activity(path)
    activity = parse_activity_file(path)
    first = activity.rows[0]
    equivalent = replace(activity, rows=[replace(first, request_code_listing=" activation ",
                                                  request=str(first.request), transaction_no=str(first.transaction_no),
                                                  authorization_code=str(first.authorization_code)), *activity.rows[1:]])
    assert activity_identity(activity) == activity_identity(equivalent)
    duplicate_row = replace(activity, rows=[*activity.rows, first])
    assert activity_identity(activity) != activity_identity(duplicate_row)
    assert identity(Decimal("-0.00")) == identity(Decimal("0.00"))


def test_harmless_numeric_serialization_keeps_workbook_identity(tmp_path):
    original = tmp_path / "original.xlsx"
    saved = tmp_path / "saved.xlsx"
    book = Workbook()
    book.active["A1"] = 50
    book.save(original)
    with ZipFile(original) as source, ZipFile(saved, "w") as destination:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                assert b"<v>50</v>" in payload
                payload = payload.replace(b"<v>50</v>", b"<v>50.0</v>")
            destination.writestr(entry, payload)
    assert workbook_values_identity(original) == workbook_values_identity(saved)


def test_same_total_correction_describes_changed_transaction_dates(tmp_path):
    paths = make_layout(tmp_path)
    source = paths["input"] / "9355 Virginia Beach" / "activity" / "week.xlsx"
    _activity(source)
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 1),
                       issue=Decimal("50"), payment=Decimal("30"))
    published = run(paths)[0]
    assert published.status == "created", published.message
    shutil.copy2(published.archive_path / "activity" / source.name, source)
    book = load_workbook(source)
    book.active["D5"] = "2026-06-06"
    book.save(source)
    correction = run(paths)[0]
    assert correction.correction_token
    assert "2026-06-07" in str(correction.changed_facts)
    assert "2026-06-06" in str(correction.changed_facts)


def test_restore_cli_returns_deliberate_correction_without_claiming_no_change(tmp_path, monkeypatch, capsys):
    from gift_card_recon import weekly_service
    def restore(**kwargs):
        raise CorrectionRequired("old-result", "new-input", {"activity_issue": "50.00 -> 51.00"})
    monkeypatch.setattr(weekly_service, "restore_weekly_revision", restore, raising=False)
    exit_code = weekly_main(["--operations-root", str(tmp_path), "--restore-revision",
                             "9355=2026-W23/r0001-0123456789abcdef", "--json"])
    records = json.loads(capsys.readouterr().out)
    assert exit_code == 1
    assert records[0]["status"] == "correction-required"
    assert records[0]["correction_token"]
    assert records[0]["run_id"]


def test_restore_cli_passes_exact_selected_revision_and_approval(tmp_path, monkeypatch, capsys):
    from gift_card_recon import weekly_service
    received = []
    def restore(**kwargs):
        received.append(kwargs)
        return WeeklyPublication(period="2026-W23", period_end=date(2026, 6, 7), status="PASS",
            output_path=tmp_path / "report.xlsx", archive_path=tmp_path / "archive",
            monthly_activity_path=tmp_path / "monthly.xlsx", variance_explanation_path=None,
            message="Restored selected prior revision", input_set_id="input", result_id="result")
    monkeypatch.setattr(weekly_service, "restore_weekly_revision", restore, raising=False)
    exit_code = weekly_main(["--operations-root", str(tmp_path), "--restore-revision",
                             "9355=2026-W23/r0001-0123456789abcdef", "--approve-correction", "9355=exact-token", "--json"])
    records = json.loads(capsys.readouterr().out)
    assert exit_code == 0
    assert records[0]["status"] == "restored"
    assert received[0]["correction_token"] == "exact-token"
    assert received[0]["revision_id"] == "r0001-0123456789abcdef"
    assert received[0]["output_root"] == tmp_path / "03 Finished Reports" / "Weekly"
    assert not list(tmp_path.iterdir())


def test_historical_duplicate_readiness_uses_archive_after_micros_export_rotates(tmp_path):
    paths = make_layout(tmp_path)
    source = paths["input"] / "9355 Virginia Beach" / "activity" / "week.xlsx"
    _activity(source)
    create_micros_week(paths["operations"], store="9355", week_start=date(2026, 6, 1),
                       issue=Decimal("50"), payment=Decimal("30"))
    published = run(paths)[0]
    assert published.status == "created", published.message
    shutil.copy2(published.archive_path / "activity" / source.name, source)
    micros = create_micros_week(paths["operations"], store="9355", week_start=date(2026, 7, 6),
                                issue=Decimal("500"), payment=Decimal("300"))
    before = _snapshot(tmp_path)
    result = assess_weekly_readiness(operations_root=paths["operations"], stores=["9355"])[0]
    assert result.ready and result.disposition == "DUPLICATE", result.reasons
    assert _snapshot(tmp_path) == before
    # One missing source is incomplete evidence, even if the remaining file is newer.
    (micros / "TENDER_DETAIL.TXT").unlink()
    before = _snapshot(tmp_path)
    partial = assess_weekly_readiness(operations_root=paths["operations"], stores=["9355"])[0]
    assert not partial.ready and partial.disposition == "BLOCKED"
    assert _snapshot(tmp_path) == before


def test_monthly_readiness_resolves_local_renamed_report_with_old_machine_paths(tmp_path, monkeypatch):
    from gift_card_recon.monthly_close_service import run_monthly_close_service
    setup = _build_period(tmp_path, store="9355")
    monkeypatch.setattr("gift_card_recon.monthly_close_service.parse_darden_credit_memo", lambda path: setup["darden"])
    published = run_monthly_close_service(
        store="9355", period="FY27-M01", input_dir=setup["input_dir"], micros_path=setup["micros_dir"],
        micros_work_dir=tmp_path / "extract", archive_root=setup["archive_root"],
        output_root=setup["output_root"], darden_report=setup["darden"], cleanup_sources=False,
        allow_unconfigured_micros=True, pdf_exporter=_fake_pdf_exporter,
    )
    renamed = published.workbook_path.with_name("Reviewed monthly close.xlsx")
    published.workbook_path.rename(renamed)
    manifest = json.loads(published.manifest_path.read_text(encoding="utf-8"))
    for artifact in manifest["artifacts"]:
        artifact["path"] = "C:/OtherComputer/" + Path(artifact["path"]).name
    published.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    before = _snapshot(tmp_path)
    result = assess_monthly_readiness(
        operations_root=tmp_path, period="FY27-M01", stores=["9355"],
        archive_root=setup["archive_root"], output_root=setup["output_root"],
        input_dirs={"9355": setup["input_dir"]}, micros_paths={"9355": setup["micros_dir"]},
        allow_unconfigured_micros=True,
    )[0]
    assert result.ready and result.disposition == "DUPLICATE", result.reasons
    assert str(renamed) in result.artifact_paths
    assert result.pdf_status == "complete"
    assert _snapshot(tmp_path) == before


def test_monthly_readiness_assesses_unstaged_intake_without_creating_period_folder(tmp_path, monkeypatch):
    setup = _build_period(tmp_path, store="9355")
    period_input = setup["input_dir"]
    store_root = period_input.parent
    setup["summary_path"].rename(store_root / setup["summary_path"].name)
    weekly = tmp_path / "01 Weekly Gift Card Activity Reports" / "9355 Virginia Beach" / "activity"
    weekly.mkdir(parents=True)
    for path in (period_input / "activity").glob("*.xlsx"):
        path.rename(weekly / path.name)
    inbox = tmp_path / "Monthly Close" / "Darden Reports - Drop Here"
    inbox.mkdir()
    darden = setup["darden"].source_file
    destination = inbox / darden.name
    darden.rename(destination)
    report = replace(setup["darden"], source_file=destination)
    for folder in (period_input / "summary", period_input / "activity", period_input / "darden", period_input):
        folder.rmdir()
    monkeypatch.setattr("gift_card_recon.monthly_close_service.parse_darden_credit_memo", lambda path: report)
    monkeypatch.setattr("gift_card_recon.monthly_close_cli.parse_darden_credit_memo", lambda path: report)
    before = _snapshot(tmp_path)
    def no_mkdir(*args, **kwargs):
        pytest.fail("Readiness attempted to create its staging destination")
    monkeypatch.setattr(Path, "mkdir", no_mkdir)
    result = assess_monthly_readiness(
        operations_root=tmp_path, period="FY27-M01", stores=["9355"], input_root=tmp_path / "Monthly Close",
        micros_paths={"9355": setup["micros_dir"]}, allow_unconfigured_micros=True,
    )[0]
    assert result.ready, result.reasons
    assert result.input_set_id
    assert not period_input.exists()
    assert _snapshot(tmp_path) == before


def test_monthly_readiness_keeps_restaurant_outcomes_independent(tmp_path, monkeypatch):
    setups = {store: _build_period(tmp_path, store=store) for store in ("9354", "9355")}
    book = load_workbook(setups["9354"]["summary_path"])
    book.active["D3"] = True
    book.save(setups["9354"]["summary_path"])
    book.close()
    reports = {str(data["darden"].source_file): data["darden"] for data in setups.values()}
    monkeypatch.setattr("gift_card_recon.monthly_close_service.parse_darden_credit_memo", lambda path: reports[str(path)])
    before = _snapshot(tmp_path)
    results = assess_monthly_readiness(
        operations_root=tmp_path, period="FY27-M01", stores=["9354", "9355"],
        input_dirs={store: data["input_dir"] for store, data in setups.items()},
        micros_paths={store: data["micros_dir"] for store, data in setups.items()},
        allow_unconfigured_micros=True,
    )
    assert [(item.store, item.ready) for item in results] == [("9354", False), ("9355", True)]
    assert "money value" in results[0].reasons[0]
    assert _snapshot(tmp_path) == before
