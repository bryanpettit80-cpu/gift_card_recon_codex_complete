from datetime import date
from decimal import Decimal
import json
from pathlib import Path
import shutil

import pytest
from openpyxl import Workbook, load_workbook

from gift_card_recon.models import WeeklyPosVariance
from gift_card_recon.monthly_close_service import _load_weekly_variance_explanations
from gift_card_recon.store_config import get_store_config
from gift_card_recon.utils import sha256_file
from gift_card_recon.variance_explanation import (
    EXPLANATION_INPUT_CELL,
    EXPLANATION_SHEET,
    WeeklyVarianceExplanation,
    add_variance_explanation_sheets,
    read_variance_explanation_workbook,
    resolve_live_weekly_explanation_path,
)


WEEK_START = date(2026, 8, 24)
WEEK_END = date(2026, 8, 30)
STORE_FOLDER = Path("9355 Virginia Beach") / "2026"
REPORT_NAME = "Gift_Card_Reconciliation_9355_2026-W35.xlsx"
EXPLANATION = "A documented payment correction explains the $35 difference."


def _weekly_package(archive_root: Path, output_root: Path) -> tuple[Path, Path]:
    package = archive_root / STORE_FOLDER / "2026-W35"
    archived = package / "report" / REPORT_NAME
    archived.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "Reconciliation"
    workbook.active["A1"] = "Synthetic weekly report"
    add_variance_explanation_sheets(workbook, WeeklyVarianceExplanation(
        "9355", WEEK_START, WEEK_END, Decimal("0"), Decimal("35"),
        Decimal("-35"), Decimal("0"),
    ))
    workbook.save(archived)
    workbook.close()
    manifest = {
        "schema_version": 1, "store": "9355", "period": "2026-W35",
        "week": {"start": WEEK_START.isoformat(), "end": WEEK_END.isoformat()},
        "variance_explanation": {"storage": "weekly_report"},
        "artifacts": {"archived_workbook": {
            "relative_path": "report/" + REPORT_NAME,
            "sha256": sha256_file(archived), "size_bytes": archived.stat().st_size,
        }, "published_workbook": {"path": "C:/historical-workspace/report.xlsx"}},
    }
    (package / "weekly_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    report = output_root / STORE_FOLDER / REPORT_NAME
    report.parent.mkdir(parents=True)
    workbook = load_workbook(archived)
    workbook[EXPLANATION_SHEET][EXPLANATION_INPUT_CELL] = EXPLANATION
    workbook.save(report)
    workbook.close()
    return package, report


@pytest.mark.parametrize("layout, explicit", [
    ("organized", False), ("organized", True),
    ("legacy", False), ("legacy", True),
    ("custom-monthly", True), ("custom-weekly", True),
])
def test_saved_weekly_explanation_resolves_supported_roots(tmp_path, layout, explicit):
    if layout == "organized":
        inputs = tmp_path / "02 Monthly Close Inputs" / "9355 Virginia Beach" / "FY27 M03"
        archive_root, output_root = tmp_path / "04 Archive", tmp_path / "03 Finished Reports"
        archive, output = archive_root / "Weekly Reconciliation", output_root / "Weekly"
    elif layout == "legacy":
        inputs = tmp_path / "Monthly Close" / "9355" / "FY27 M03"
        archive_root, output_root = tmp_path / "Archive - Old Files", tmp_path / "Output"
        archive, output = archive_root / "Weekly Reconciliation", output_root
    else:
        inputs = tmp_path / "arbitrary-inputs"
        archive_root, output_root = tmp_path / "custom-evidence", tmp_path / "custom-reports"
        archive = archive_root / "Weekly Reconciliation" if layout == "custom-monthly" else archive_root
        output = output_root / "Weekly" if layout == "custom-monthly" else output_root
    _, report = _weekly_package(archive, output)
    kwargs = {"archive_root": archive_root, "output_root": output_root} if explicit else {}
    resolved = resolve_live_weekly_explanation_path(inputs, "9355", WEEK_END, **kwargs)
    assert resolved == report.resolve()
    assert read_variance_explanation_workbook(resolved).explanation == EXPLANATION


@pytest.mark.parametrize("weekly_roots", [False, True])
def test_monthly_loader_uses_explicit_roots_for_embedded_explanations(tmp_path, weekly_roots):
    archive_root, output_root = tmp_path / "chosen-archive", tmp_path / "chosen-output"
    archive = archive_root if weekly_roots else archive_root / "Weekly Reconciliation"
    output = output_root if weekly_roots else output_root / "Weekly"
    _, report = _weekly_package(archive, output)
    variance = WeeklyPosVariance(
        week_ending=WEEK_END, report_begin=WEEK_START, report_end=WEEK_END,
        activity_issue=Decimal("100"), pos_issue=Decimal("100"), issue_variance=Decimal("0"),
        activity_payment=Decimal("200"), pos_payment=Decimal("235"), payment_variance=Decimal("35"),
        net_variance=Decimal("-35"), coverage_status="PASS",
    )
    explanations, paths, hashes, blockers = _load_weekly_variance_explanations(
        input_dir=tmp_path / "custom-inputs", config=get_store_config("9355"),
        weekly_variances=[variance], weekly_tender={"Week ending 08/30/2026 tender": Decimal("0")},
        archived_variance_explanations=None, archive_root=archive_root, output_root=output_root,
    )
    assert explanations[WEEK_END].explanation == EXPLANATION
    assert paths == (report.resolve(),)
    assert hashes[report.resolve()] == sha256_file(report)
    assert not blockers


def test_explicit_roots_do_not_fall_back_to_another_operations_tree(tmp_path):
    inputs = tmp_path / "02 Monthly Close Inputs" / "9355 Virginia Beach" / "FY27 M03"
    _weekly_package(tmp_path / "04 Archive" / "Weekly Reconciliation", tmp_path / "03 Finished Reports" / "Weekly")
    assert resolve_live_weekly_explanation_path(
        inputs, "9355", WEEK_END,
        archive_root=tmp_path / "different-archive", output_root=tmp_path / "different-output",
    ) is None


@pytest.mark.parametrize("duplicate", ["archive", "report"])
def test_ambiguous_configured_weekly_layout_is_rejected(tmp_path, duplicate):
    archive, output = tmp_path / "archive", tmp_path / "output"
    package, report = _weekly_package(archive / "Weekly Reconciliation", output / "Weekly")
    if duplicate == "archive":
        shutil.copytree(package, archive / STORE_FOLDER / "2026-W35")
        message = "Multiple weekly explanation archive packages"
    else:
        other = output / STORE_FOLDER / REPORT_NAME
        other.parent.mkdir(parents=True)
        shutil.copy2(report, other)
        message = "Multiple editable weekly reports"
    with pytest.raises(ValueError, match=message):
        resolve_live_weekly_explanation_path(
            tmp_path / "inputs", "9355", WEEK_END, archive_root=archive, output_root=output,
        )


@pytest.mark.parametrize("damage", ["manifest-identity", "archive-bytes", "report-values"])
def test_custom_roots_retain_identity_and_integrity_checks(tmp_path, damage):
    archive, output = tmp_path / "archive", tmp_path / "output"
    package, report = _weekly_package(archive, output)
    if damage == "manifest-identity":
        manifest_path = package / "weekly_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["store"] = "9354"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        message = "identity mismatch"
    elif damage == "archive-bytes":
        with (package / "report" / REPORT_NAME).open("ab") as handle:
            handle.write(b"changed")
        message = "integrity check failed"
    else:
        workbook = load_workbook(report)
        workbook["Reconciliation"]["A1"] = "Changed outside the explanation box"
        workbook.save(report)
        workbook.close()
        message = "values or formulas changed outside"
    with pytest.raises(ValueError, match=message):
        resolve_live_weekly_explanation_path(
            tmp_path / "inputs", "9355", WEEK_END, archive_root=archive, output_root=output,
        )


@pytest.mark.parametrize(
    "damage",
    [
        "top-level-array",
        "variance-array",
        "week-array",
        "missing-artifacts",
        "archived-record-null",
        "archived-record-missing-fields",
    ],
)
def test_malformed_weekly_manifest_shape_is_a_contextual_validation_error(tmp_path, damage):
    archive, output = tmp_path / "archive", tmp_path / "output"
    package, _ = _weekly_package(archive, output)
    manifest_path = package / "weekly_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if damage == "top-level-array":
        manifest = []
    elif damage == "variance-array":
        manifest["variance_explanation"] = []
    elif damage == "week-array":
        manifest["week"] = []
    elif damage == "missing-artifacts":
        manifest.pop("artifacts")
    elif damage == "archived-record-null":
        manifest["artifacts"]["archived_workbook"] = None
    else:
        manifest["artifacts"]["archived_workbook"] = {}
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    with pytest.raises(ValueError, match="Weekly explanation manifest structure is invalid"):
        resolve_live_weekly_explanation_path(
            tmp_path / "inputs", "9355", WEEK_END, archive_root=archive, output_root=output,
        )


def test_malformed_weekly_revision_shape_is_a_contextual_validation_error(tmp_path):
    archive, output = tmp_path / "archive", tmp_path / "output"
    package, _ = _weekly_package(archive, output)
    (package / "weekly_report_revision.json").write_text("[]", encoding="utf-8")

    with pytest.raises(ValueError, match="Weekly report revision manifest structure is invalid"):
        resolve_live_weekly_explanation_path(
            tmp_path / "inputs", "9355", WEEK_END, archive_root=archive, output_root=output,
        )


def test_non_integer_weekly_revision_size_is_a_contextual_validation_error(tmp_path):
    archive, output = tmp_path / "archive", tmp_path / "output"
    package, _ = _weekly_package(archive, output)
    manifest_path = package / "weekly_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    revision = {
        "schema_version": 1,
        "store": manifest["store"],
        "period": manifest["period"],
        "week": manifest["week"],
        "original_manifest_sha256": sha256_file(manifest_path),
        "original_report_sha256": manifest["artifacts"]["archived_workbook"]["sha256"],
        "revised_workbook": {
            "relative_path": f"report-revisions/revision-1/{REPORT_NAME}",
            "size_bytes": float("inf"),
            "sha256": "0" * 64,
        },
    }
    (package / "weekly_report_revision.json").write_text(json.dumps(revision), encoding="utf-8")

    with pytest.raises(ValueError, match="Weekly report revision manifest structure is invalid"):
        resolve_live_weekly_explanation_path(
            tmp_path / "inputs", "9355", WEEK_END, archive_root=archive, output_root=output,
        )
