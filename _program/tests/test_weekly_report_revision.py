from datetime import date
from decimal import Decimal
import json
from pathlib import Path
import shutil

import pytest
from openpyxl import Workbook, load_workbook

from gift_card_recon import weekly_report_revision
from gift_card_recon.utils import sha256_file
from gift_card_recon.variance_explanation import WeeklyVarianceExplanation, add_variance_explanation_sheets
from gift_card_recon.weekly_report_revision import publish_weekly_report_revision, resolve_weekly_report_revision


@pytest.fixture
def revision_inputs(tmp_path: Path):
    package = tmp_path / "archive" / "2026-W35"
    original = package / "report" / "Gift_Card_Reconciliation_9355_2026-W35.xlsx"
    original.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "Reconciliation"
    workbook.active["C6"] = 870
    workbook.active["A1"] = "Weekly gift card reconciliation"
    workbook.save(original)
    workbook.close()
    canonical = tmp_path / "weekly" / original.name
    canonical.parent.mkdir()
    shutil.copy2(original, canonical)
    manifest = {
        "schema_version": 1, "store": "9355", "period": "2026-W35",
        "week": {"start": "2026-08-24", "end": "2026-08-30"},
        "artifacts": {"archived_workbook": {
            "relative_path": "report/" + original.name,
            "sha256": sha256_file(original), "size_bytes": original.stat().st_size,
        }},
    }
    (package / "weekly_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    candidate = tmp_path / "candidate.xlsx"
    workbook = load_workbook(canonical)
    workbook.active["C6"] = 920
    workbook.security = None  # Excel-saved workbooks can omit workbook protection.
    add_variance_explanation_sheets(workbook, WeeklyVarianceExplanation(
        "9355", date(2026, 8, 24), date(2026, 8, 30), Decimal("0"), Decimal("35"),
        Decimal("-35"), Decimal("0"), "The reload is included in issuance; payment follow-up remains.",
    ))
    workbook.save(candidate)
    workbook.close()
    return package, original, canonical, candidate


def test_report_revision_retains_original_evidence_and_has_verified_baseline(revision_inputs):
    package, original, canonical, candidate = revision_inputs
    original_hash = sha256_file(original)
    manifest_bytes = (package / "weekly_manifest.json").read_bytes()
    sidecar = publish_weekly_report_revision(package, canonical, candidate)
    assert sidecar.is_file()
    assert sha256_file(original) == original_hash
    assert (package / "weekly_manifest.json").read_bytes() == manifest_bytes
    assert sha256_file(canonical) == sha256_file(candidate)
    baseline = resolve_weekly_report_revision(package)
    assert baseline is not None
    assert baseline.parent.parent.name == "report-revisions"
    assert sha256_file(baseline) == sha256_file(candidate)


def test_report_revision_accepts_reviewed_presentation_only_changes(revision_inputs):
    package, original, canonical, candidate = revision_inputs
    workbook = load_workbook(canonical)
    workbook.active.column_dimensions["C"].width = 35
    workbook.save(canonical)
    workbook.close()
    current_hash = sha256_file(canonical)
    assert current_hash != sha256_file(original)
    with pytest.raises(ValueError, match="differs from its original archive"):
        publish_weekly_report_revision(package, canonical, candidate)
    publish_weekly_report_revision(package, canonical, candidate, expected_canonical_sha256=current_hash)
    assert resolve_weekly_report_revision(package).is_file()


def test_report_revision_refuses_changed_values_even_with_reviewed_hash(revision_inputs):
    package, original, canonical, candidate = revision_inputs
    workbook = load_workbook(canonical)
    workbook.active["C6"] = 999
    workbook.save(canonical)
    workbook.close()
    current_hash = sha256_file(canonical)
    with pytest.raises(ValueError, match="values or formulas changed"):
        publish_weekly_report_revision(package, canonical, candidate, expected_canonical_sha256=current_hash)
    assert sha256_file(canonical) == current_hash
    assert not (package / "weekly_report_revision.json").exists()


def test_report_revision_rolls_back_current_presentation_when_sidecar_publish_fails(
    revision_inputs, monkeypatch,
):
    package, original, canonical, candidate = revision_inputs
    workbook = load_workbook(canonical)
    workbook.active.column_dimensions["C"].width = 35
    workbook.save(canonical)
    workbook.close()
    prior_hash = sha256_file(canonical)
    real_replace = weekly_report_revision.os.replace
    def fail_sidecar(source, destination):
        if Path(destination).name == "weekly_report_revision.json":
            raise PermissionError("Simulated Dropbox lock")
        return real_replace(source, destination)
    monkeypatch.setattr(weekly_report_revision.os, "replace", fail_sidecar)
    with pytest.raises(PermissionError, match="Dropbox lock"):
        publish_weekly_report_revision(package, canonical, candidate, expected_canonical_sha256=prior_hash)
    assert sha256_file(canonical) == prior_hash
    assert not (package / "weekly_report_revision.json").exists()
    assert not list((package / "report-revisions").rglob("*.xlsx"))


def test_report_revision_rejects_tampered_baseline_and_path_traversal(revision_inputs):
    package, original, canonical, candidate = revision_inputs
    sidecar = publish_weekly_report_revision(package, canonical, candidate)
    baseline = resolve_weekly_report_revision(package)
    baseline.write_bytes(b"changed evidence")
    with pytest.raises(ValueError, match="integrity check failed"):
        resolve_weekly_report_revision(package)
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    payload["revised_workbook"]["relative_path"] = "report-revisions/../" + original.name
    sidecar.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError, match="outside its revision folder"):
        resolve_weekly_report_revision(package)
